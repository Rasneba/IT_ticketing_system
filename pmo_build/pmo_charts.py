"""Chart helpers for the Assistant PMO dashboards."""

from openpyxl.chart import BarChart, DoughnutChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.marker import Marker

try:
    from openpyxl.drawing.text import ParagraphProperties, CharacterProperties
    _HAS_TEXT = True
except Exception:  # pragma: no cover
    _HAS_TEXT = False

NAVY = "1F4E79"
GREEN = "2E7D32"
AMBER = "E8A33D"
RED = "C0392B"
DARKRED = "8E2C22"
BLUE = "2E75B6"
BLUE_LT = "7FA9D4"
GRAY = "9AA7B4"


def _title(chart, text, size=1200, color=NAVY):
    chart.title = text
    if _HAS_TEXT:
        try:
            chart.title.tx.rich.p[0].pPr = ParagraphProperties(
                defRPr=CharacterProperties(sz=size, b=True, solidFill=color))
        except Exception:
            pass


def _axis_text(chart, size=900, color="44515F"):
    if not _HAS_TEXT:
        return
    for ax in (chart.x_axis, chart.y_axis):
        try:
            ax.txPr = None
            ax.title = ax.title
        except Exception:
            pass


def _point_colors(series, colors):
    points = []
    for i, c in enumerate(colors):
        gp = GraphicalProperties(solidFill=c)
        gp.line.solidFill = "FFFFFF"
        gp.line.width = 12700
        points.append(DataPoint(idx=i, spPr=gp))
    series.dPt = points


def _solid(series, color, border="FFFFFF"):
    gp = GraphicalProperties(solidFill=color)
    if border:
        gp.line.solidFill = border
        gp.line.width = 12700
    series.graphicalProperties = gp


def doughnut(ws, title, cats, vals, colors, anchor, width=17.5, height=7.6,
             labels="percent", legend="b", hole=55):
    ch = DoughnutChart()
    ch.add_data(vals, titles_from_data=True)
    ch.set_categories(cats)
    try:
        ch.holeSize = hole
    except Exception:
        pass
    if colors:
        _point_colors(ch.series[0], colors)
    _title(ch, title)
    if labels == "percent":
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showPercent = True
        ch.dataLabels.showVal = False
        ch.dataLabels.showCatName = False
        ch.dataLabels.showSerName = False
    elif labels == "val":
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
    if legend:
        ch.legend.position = legend
        ch.legend.overlay = False
    else:
        ch.legend = None
    ch.height, ch.width = height, width
    ws.add_chart(ch, anchor)
    return ch


def bar(ws, title, cats, data, anchor, colors=None, horizontal=False, width=17.5,
        height=7.6, labels=True, legend=True, gap=60, num_fmt=None, point_colors=None):
    ch = BarChart()
    ch.type = "bar" if horizontal else "col"
    ch.grouping = "clustered"
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    for i, s in enumerate(ch.series):
        if point_colors and i == 0:
            _point_colors(s, point_colors)
        elif colors:
            if isinstance(colors, str):
                _solid(s, colors)
            else:
                _solid(s, colors[i % len(colors)])
        else:
            _solid(s, BLUE)
    _title(ch, title)
    if labels:
        ch.dataLabels = DataLabelList()
        ch.dataLabels.showVal = True
        if num_fmt:
            ch.dataLabels.numFmt = num_fmt
    if legend and len(ch.series) > 1:
        ch.legend.position = "b"
        ch.legend.overlay = False
    else:
        ch.legend = None
    ch.gapWidth = gap
    ch.height, ch.width = height, width
    try:
        ch.x_axis.delete = False
        ch.y_axis.delete = False
    except Exception:
        pass
    ws.add_chart(ch, anchor)
    return ch


def line(ws, title, cats, data, anchor, colors=None, width=17.5, height=7.6, legend=True):
    ch = LineChart()
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    for i, s in enumerate(ch.series):
        c = (colors or [BLUE, GREEN])[i % len(colors or [BLUE, GREEN])]
        s.graphicalProperties = GraphicalProperties()
        s.graphicalProperties.line.solidFill = c
        s.graphicalProperties.line.width = 22000
        s.marker = Marker(symbol="circle", size=6)
        s.marker.graphicalProperties = GraphicalProperties(solidFill=c)
        s.marker.graphicalProperties.line.solidFill = c
        s.smooth = False
    _title(ch, title)
    if legend:
        ch.legend.position = "b"
        ch.legend.overlay = False
    else:
        ch.legend = None
    ch.height, ch.width = height, width
    try:
        ch.x_axis.delete = False
        ch.y_axis.delete = False
    except Exception:
        pass
    ws.add_chart(ch, anchor)
    return ch
