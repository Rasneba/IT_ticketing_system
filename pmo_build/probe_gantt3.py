"""Parameterised copy of add_gantt to isolate the Excel-rejected element."""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Side, Border
from openpyxl.utils import get_column_letter
import win32com.client as win32

import build as B
import pmo_specs as S1
from pmo_core import (NAVY, BLUE, BLUE_MID, BLUE_LT, WHITE, BORDER_C,
                      G_BG, G_TX, A_BG, A_TX, R_BG, R_TX, N_BG, N_TX,
                      font, solid, al, band, setup_print, cf_formula, col_letter_map)

TEMP = os.path.join(os.environ["TEMP"], "pmo_g4.xlsx")
customers = B.load_customers()
FIRST, LAST = 5, 34


def gantt(ws, weeks=True, row4=True, grid=True, cf_status=True,
          cf_catch=True, cf_week=True, band3=True, band36=True, printing=True):
    spec = S1.WBS_SPEC
    cols = spec["cols"]
    L = col_letter_map(cols)
    ncols = len(cols)
    W = S1.GANTT_WEEKS
    g1 = ncols + 1
    g2 = ncols + W
    L1, L2 = get_column_letter(g1), get_column_letter(g2)
    start_c, fin_c, st_c = L["Start Date"], L["Finish Date"], L["Status"]

    if band3:
        band(ws, 3, g2, "GANTT", BLUE)
    ws.row_dimensions[3].height = 18

    if weeks:
        for i in range(W):
            col = g1 + i
            cl = get_column_letter(col)
            cell = ws.cell(row=4, column=col)
            if row4:
                cell.value = "='23_SETTINGS'!$B$5" if i == 0 else f"={get_column_letter(col - 1)}4+7"
            cell.number_format = "dd-mmm"
            cell.font = font(8, True, NAVY)
            cell.alignment = Alignment(textRotation=90, horizontal="center", vertical="bottom")
            cell.fill = solid(BLUE_LT)
            cell.border = Border(left=Side(style="hair", color=BORDER_C),
                                 right=Side(style="hair", color=BORDER_C),
                                 top=Side(style="hair", color=BORDER_C),
                                 bottom=Side(style="medium", color=BLUE_MID))
            ws.column_dimensions[cl].width = 3.6
    ws.row_dimensions[4].height = 56

    gridb = Border(*[Side(style="hair", color=BORDER_C)] * 4)
    if grid:
        for r in range(FIRST, LAST + 1):
            for i in range(W):
                col = g1 + i
                cl = get_column_letter(col)
                cell = ws.cell(row=r, column=col)
                cell.value = (f'=IF(AND(${start_c}{r}<>"",${fin_c}{r}<>"",'
                              f'${start_c}{r}<={cl}$4+6,${fin_c}{r}>={cl}$4),1,"")')
                cell.number_format = "General"
                cell.alignment = al("center", "center")
                cell.font = font(8)
                cell.fill = solid(WHITE)
                cell.border = gridb

    rng = f"{L1}{FIRST}:{L2}{LAST}"
    if cf_status:
        for status, bg, tx in (("Complete", BLUE, WHITE), ("In Progress", BLUE_MID, WHITE),
                               ("Blocked", R_BG, R_TX)):
            cf_formula(ws, rng, f'AND(ISNUMBER({L1}{FIRST}),${st_c}{FIRST}="{status}")',
                       bg, tx, first_row=FIRST, stop=True)
    if cf_catch:
        cf_formula(ws, rng, f"ISNUMBER({L1}{FIRST})", N_BG, N_TX, first_row=FIRST, stop=True)
    if cf_week:
        cf_formula(ws, f"{L1}4:{L2}4", f"AND({L1}$4<=TODAY(),{L1}$4+6>=TODAY())",
                   R_BG, R_TX, bold=True, first_row=4, stop=True)

    if band36:
        band(ws, 36, g2, "LEGEND", BLUE)
        ws.row_dimensions[36].height = 18

    if printing:
        setup_print(ws, g2, 36, title_rows="1:4")


def run(tag, **kw):
    wb = Workbook()
    wb.remove(wb.active)
    for spec in B.SPECS:
        B.build_table_sheet(wb, spec)
    B.build_settings(wb, customers)
    gantt(wb["04_WBS_SCHEDULE"], **kw)
    wb.calculation.fullCalcOnLoad = True
    wb.save(TEMP)
    try:
        bk = xl.Workbooks.Open(TEMP)
        print(f"OK   {tag}")
        bk.Close(False)
    except Exception:
        print(f"FAIL {tag}")


xl = win32.Dispatch("Excel.Application")
xl.Visible = False
xl.DisplayAlerts = False
try:
    run("full")
    run("no-weeks", weeks=False)
    run("no-row4", row4=False)
    run("no-grid", grid=False)
    run("no-cf-status", cf_status=False)
    run("no-cf-catch", cf_catch=False)
    run("no-cf-week", cf_week=False)
    run("no-band3", band3=False)
    run("no-band36", band36=False)
    run("no-print", printing=False)
    run("bare", weeks=False, grid=False, cf_status=False, cf_catch=False,
        cf_week=False, band3=False, band36=False, printing=False)
finally:
    xl.Quit()
