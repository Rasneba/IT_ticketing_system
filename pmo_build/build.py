"""Build Assistant_PMO_Production_PMOSystem.xlsx.

Run:  python pmo_build\\build.py
"""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

import pmo_specs as S1
import pmo_specs2 as S2
from pmo_core import build_table_sheet
from pmo_forms import build_home, build_charter, build_weekly_status, build_sop
from pmo_settings import build_settings
from pmo_extras import add_gantt, add_summary_block, add_quadrant, build_data_dictionary
from pmo_dash_exec import build_exec_dashboard
from pmo_dash_proj import build_project_dashboard

OUT = os.path.join(ROOT, "Assistant_PMO_Production_PMOSystem.xlsx")
CUSTOMERS_FILE = os.path.join(ROOT, "V7 Excel sheet  (6).xlsx")

SPECS = [
    S1.PORTFOLIO_SPEC, S1.WBS_SPEC, S1.MILESTONE_SPEC, S1.RAID_SPEC, S1.ACTION_SPEC,
    S2.RACI_SPEC, S2.STAKEHOLDER_SPEC, S2.CHANGE_SPEC, S2.COST_SPEC, S2.RESOURCE_SPEC,
    S2.PROCUREMENT_SPEC, S2.QUALITY_SPEC, S2.MEETING_SPEC, S2.DECISION_SPEC, S2.COMM_SPEC,
    S2.CLOSURE_SPEC, S2.LESSON_SPEC, S2.PLAN_SPEC,
]

ORDER = [
    "00_HOME",
    "01_EXECUTIVE_DASHBOARD",
    "01B_PROJECT_DASHBOARD",
    "02_PROJECT_PORTFOLIO",
    "03_PROJECT_CHARTER",
    "04_WBS_SCHEDULE",
    "05_MILESTONES",
    "06_RAID_LOG",
    "07_ACTION_TRACKER",
    "08_RACI_MATRIX",
    "09_STAKEHOLDERS",
    "10_CHANGE_CONTROL",
    "11_COST_MANAGEMENT",
    "12_RESOURCE_MANAGEMENT",
    "13_PROCUREMENT",
    "14_QUALITY_CONTROL",
    "15_WEEKLY_STATUS",
    "16_MEETING_MINUTES",
    "17_DECISION_LOG",
    "18_COMMUNICATION_PLAN",
    "19_PROJECT_CLOSURE",
    "20_LESSONS_LEARNED",
    "21_PMOSOP",
    "22_30_60_90_PLAN",
    "23_SETTINGS",
    "24_DATA_DICTIONARY",
]


def load_customers():
    wb = load_workbook(CUSTOMERS_FILE, read_only=True, data_only=True)
    ws = wb["ALL CUSTOMER"]
    seen, out = set(), ["Internal"]
    seen.add("internal")
    for (value,) in ws.iter_rows(min_row=2, min_col=2, max_col=2, values_only=True):
        if not value:
            continue
        text = str(value).strip()
        if not text or len(text) > 60 or text.lower() in seen:
            continue
        seen.add(text.lower())
        out.append(text)
    wb.close()
    out[1:] = sorted(out[1:], key=str.lower)
    return out


def main():
    customers = load_customers()
    print(f"customers loaded: {len(customers)}")

    wb = Workbook()
    wb.remove(wb.active)

    for spec in SPECS:
        build_table_sheet(wb, spec)
    print(f"table sheets: {len(SPECS)}")

    build_home(wb)
    build_exec_dashboard(wb)
    build_project_dashboard(wb)
    build_charter(wb)
    build_weekly_status(wb)
    build_sop(wb)
    build_settings(wb, customers)
    build_data_dictionary(wb)
    print("special sheets: 8")

    add_gantt(wb["04_WBS_SCHEDULE"])
    add_summary_block(wb["11_COST_MANAGEMENT"], "cost")
    add_summary_block(wb["14_QUALITY_CONTROL"], "quality")
    add_summary_block(wb["19_PROJECT_CLOSURE"], "closure")
    add_quadrant(wb["09_STAKEHOLDERS"])
    print("post blocks: gantt / cost / quality / closure / matrix")

    wb.defined_names["L_Projects"] = DefinedName(
        "L_Projects", attr_text="'02_PROJECT_PORTFOLIO'!$C$5:$C$34")

    missing = [n for n in ORDER if n not in wb.sheetnames]
    extra = [n for n in wb.sheetnames if n not in ORDER]
    if missing or extra:
        raise SystemExit(f"sheet mismatch: missing={missing} extra={extra}")
    wb._sheets.sort(key=lambda ws: ORDER.index(ws.title))
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True

    wb.save(OUT)

    charts = sum(len(ws._charts) for ws in wb.worksheets)
    tables = sum(len(ws.tables) for ws in wb.worksheets)
    dvs = sum(len(ws.data_validations.dataValidation) for ws in wb.worksheets)
    cfs = sum(len(ws.conditional_formatting._cf_rules) for ws in wb.worksheets)
    print(f"saved: {OUT}")
    print(f"sheets={len(wb.worksheets)} tables={tables} dv={dvs} cf_ranges={cfs} charts={charts}"
          f" defined_names={len(wb.defined_names)}")


if __name__ == "__main__":
    main()
