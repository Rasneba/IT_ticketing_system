"""Isolate the gantt failure by mutating the built sheet."""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
import win32com.client as win32
import build as B

TEMP = os.path.join(os.environ["TEMP"], "pmo_gantt.xlsx")
customers = B.load_customers()


def base():
    wb = Workbook()
    wb.remove(wb.active)
    for spec in B.SPECS:
        B.build_table_sheet(wb, spec)
    B.build_settings(wb, customers)
    return wb


def run(tag, mutate):
    wb = base()
    ws = wb["04_WBS_SCHEDULE"]
    B.add_gantt(ws)
    mutate(ws)
    wb.calculation.fullCalcOnLoad = True
    wb.save(TEMP)
    try:
        bk = xl.Workbooks.Open(TEMP)
        print(f"OK   {tag}")
        bk.Close(False)
    except Exception:
        print(f"FAIL {tag}")


def noop(ws):
    pass


def clear_cf(ws):
    ws.conditional_formatting._cf_rules = {}


def clear_grid(ws):
    n = len(S1.WBS_SPEC["cols"])
    for r in range(5, 35):
        for c in range(n + 1, n + 25):
            ws.cell(row=r, column=c).value = None


def clear_week_cf(ws):
    n = len(S1.WBS_SPEC["cols"])
    from openpyxl.utils import get_column_letter as gl
    rng = f"{gl(n+1)}4:{gl(n+24)}4"
    key = None
    for k in list(ws.conditional_formatting._cf_rules.keys()):
        if str(k.sqref) == rng:
            key = k
    if key is not None:
        del ws.conditional_formatting._cf_rules[key]


import pmo_specs as S1

xl = win32.Dispatch("Excel.Application")
xl.Visible = False
xl.DisplayAlerts = False
try:
    run("full", noop)
    run("no-cf", clear_cf)
    run("no-grid-formulas", clear_grid)
    run("no-week-cf", clear_week_cf)
finally:
    xl.Quit()
