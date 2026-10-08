"""Core styles, helpers and generic sheet builders for the Assistant PMO workbook."""

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.page import PageMargins

FONT = "Calibri"
NAVY = "0F2A44"
BLUE = "1F4E79"
BLUE_MID = "2E75B6"
BLUE_LT = "DCE9F5"
PALE = "F2F7FC"
WHITE = "FFFFFF"
GRAY_TX = "5B6B7B"
BORDER_C = "C7D2DC"

G_BG, G_TX = "D9F0DC", "17602A"
A_BG, A_TX = "FFF1CC", "8A5A00"
R_BG, R_TX = "FBDAD7", "96201D"
N_BG, N_TX = "E4E9EF", "3B4A59"

GREEN = "2E7D32"
AMBER = "E8A33D"
RED = "C0392B"
DARKRED = "8E2C22"

SAMPLE_NOTE = ("SAMPLE DATA - every sample record is marked SAMPLE in the Notes column. "
               "Delete or overwrite them freely: all formulas, KPIs, charts and drop-downs keep working.")

FMT = {
    "text": "@",
    "wrap": "@",
    "date": "dd-mmm-yy",
    "money": "#,##0",
    "pct": "0%",
    "int": "0",
    "num": "0",
    "calc": None,
}


def solid(color):
    return PatternFill("solid", fgColor=color)


def border_all(color=BORDER_C, style="thin"):
    s = Side(style=style, color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def no_border():
    return Border()


def font(size=10, bold=False, color="1F2933", italic=False, name=FONT):
    return Font(name=name, size=size, bold=bold, color=color, italic=italic)


def al(h="left", v="center", wrap=False, indent=0):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent)


# --------------------------------------------------------------------------- #
# sheet chrome
# --------------------------------------------------------------------------- #
def sheet_title(ws, title, purpose, ncols, note=None, subtitle_formula=None):
    """Rows 1-3: title band / purpose band / note band."""
    last = get_column_letter(ncols)
    ws.merge_cells(f"A1:{last}1")
    c = ws["A1"]
    c.value = title
    c.font = font(16, True, WHITE)
    c.fill = solid(NAVY)
    c.alignment = al("left", "center", indent=1)
    ws.row_dimensions[1].height = 32

    ws.merge_cells(f"A2:{last}2")
    c = ws["A2"]
    c.value = subtitle_formula if subtitle_formula else purpose
    c.font = font(9.5, False, WHITE, italic=True)
    c.fill = solid(BLUE)
    c.alignment = al("left", "center", indent=1)
    ws.row_dimensions[2].height = 17

    ws.merge_cells(f"A3:{last}3")
    c = ws["A3"]
    c.value = note if note else purpose
    c.font = font(8.5, False, GRAY_TX, italic=True)
    c.fill = solid(PALE)
    c.alignment = al("left", "center", indent=1)
    ws.row_dimensions[3].height = 15


def unmerge_overlap(ws, r1, c1, r2, c2):
    """Drop any existing merged range that intersects the target rectangle."""
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row <= r2 and rng.max_row >= r1 and rng.min_col <= c2 and rng.max_col >= c1:
            ws.merged_cells.remove(rng)


def band(ws, row, ncols, text, color=BLUE_MID, size=10):
    last = get_column_letter(ncols)
    unmerge_overlap(ws, row, 1, row, ncols)
    ws.merge_cells(f"A{row}:{last}{row}")
    c = ws[f"A{row}"]
    c.value = text
    c.font = font(size, True, WHITE)
    c.fill = solid(color)
    c.alignment = al("left", "center", indent=1)
    ws.row_dimensions[row].height = 20


def write_header(ws, cols, row=4, height=32):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c["h"])
        cell.font = font(9.5, True, WHITE)
        cell.fill = solid(NAVY)
        cell.alignment = al("center", "center", wrap=True)
        cell.border = Border(bottom=Side(style="medium", color=BLUE_MID))
    ws.row_dimensions[row].height = height


def setup_print(ws, ncols, last_row, title_rows="1:4", landscape=True, footer_left="Assistant PMO Management System"):
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.45, bottom=0.45, header=0.2, footer=0.25)
    ws.print_title_rows = title_rows
    ws.print_area = f"A1:{get_column_letter(ncols)}{last_row}"
    ws.oddFooter.left.text = footer_left
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.color = "808080"
    ws.oddFooter.right.text = "&A  |  Page &P of &N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.color = "808080"


# --------------------------------------------------------------------------- #
# data validation helpers
# --------------------------------------------------------------------------- #
def add_list_validation(ws, col_letter, list_name, first_row, last_row, prompt=None, first_col_letter=None):
    rng = f"{col_letter}{first_row}:{(first_col_letter or col_letter)}{last_row}"
    dv = DataValidation(
        type="list",
        formula1=list_name,
        allow_blank=True,
        showErrorMessage=True,
        showInputMessage=True,
        errorStyle="stop",
        errorTitle="Value not allowed",
        error=f"Pick a value from the drop-down list ({list_name}).",
        promptTitle="Select",
        prompt=prompt or "Choose a value from the drop-down list.",
    )
    dv.add(rng)
    ws.add_data_validation(dv)
    return dv


# --------------------------------------------------------------------------- #
# conditional formatting helpers
# --------------------------------------------------------------------------- #
def cf_map(ws, rng, col_letter, mapping, first_row=5, stop=True):
    """mapping: value -> (background, text colour)."""
    for value, (bg, tx) in mapping.items():
        rule = FormulaRule(
            formula=[f'EXACT(${col_letter}{first_row},"{value}")'],
            fill=solid(bg),
            font=font(10, True, tx),
            stopIfTrue=stop,
        )
        ws.conditional_formatting.add(rng, rule)


def cf_formula(ws, rng, formula, bg, tx="1F2933", bold=False, first_row=5, stop=False):
    if formula.startswith("="):
        formula = formula[1:]
    rule = FormulaRule(formula=[formula], fill=solid(bg), font=font(10, bold, tx), stopIfTrue=stop)
    ws.conditional_formatting.add(rng, rule)


def cf_databar(ws, rng, color=BLUE_MID):
    ws.conditional_formatting.add(
        rng, DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=color)
    )


RAG_MAP = {"Green": (G_BG, G_TX), "Yellow": (A_BG, A_TX), "Red": (R_BG, R_TX)}
STATUS_MAP = {
    "Active": (G_BG, G_TX),
    "Completed": (N_BG, N_TX),
    "On Hold": (A_BG, A_TX),
    "Not Started": (N_BG, N_TX),
    "Cancelled": (R_BG, R_TX),
}


# --------------------------------------------------------------------------- #
# generic table sheet builder
# --------------------------------------------------------------------------- #
def col_letter_map(cols):
    return {c["h"]: get_column_letter(i) for i, c in enumerate(cols, 1)}


def build_table_sheet(wb, spec):
    """Create one PMO sheet: title bands, header, table, values, formulas, DV, CF, print."""
    name = spec["sheet"]
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = spec.get("tab", BLUE_MID)
    ws.sheet_view.showGridLines = False

    cols = spec["cols"]
    ncols = len(cols)
    cap = spec.get("capacity", 30)
    first_row, last_row = 5, 4 + cap
    L = col_letter_map(cols)

    sheet_title(ws, spec["title"], spec["purpose"], ncols, spec.get("note", SAMPLE_NOTE))

    for i, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(i)].width = c["w"]

    write_header(ws, cols)

    wrap_used = any(c["t"] == "wrap" for c in cols)
    row_h = spec.get("row_h", 28 if wrap_used else 15)
    fmt_default = {"date": FMT["date"], "money": FMT["money"], "pct": FMT["pct"], "int": FMT["int"],
                   "num": FMT["num"], "wrap": FMT["wrap"], "text": FMT["text"]}

    samples = spec.get("sample", [])
    sample_map = {i: s for i, s in enumerate(samples)}  # row index (0-based) -> dict

    for r in range(first_row, last_row + 1):
        ws.row_dimensions[r].height = row_h
        vals = sample_map.get(r - first_row, {})
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=r, column=i)
            if c.get("f"):
                cell.value = c["f"](L, r)
            elif c["h"] in vals:
                cell.value = vals[c["h"]]
            # number format
            fmt = c.get("fmt") or fmt_default.get(c["t"])
            if fmt:
                cell.number_format = fmt
            # alignment / font
            t = c["t"]
            if t == "wrap":
                cell.alignment = al("left", "top", wrap=True)
            elif t == "date":
                cell.alignment = al("center", "center")
            elif t in ("money", "pct", "int", "num"):
                cell.alignment = al("right", "center")
            else:
                cell.alignment = al("left", "center")
            cell.font = font(10, False, "1F2933")

    # table
    last_col = get_column_letter(ncols)
    tab = Table(displayName=spec["table"], ref=f"A4:{last_col}{last_row}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                        showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(tab)

    # validations
    for c in cols:
        if c.get("lst"):
            add_list_validation(ws, get_column_letter(cols.index(c) + 1), c["lst"], first_row, last_row,
                                prompt=c.get("prompt"))

    # conditional formatting
    for rule in spec.get("cf", []):
        kind = rule[0]
        if kind == "map":
            _, header, mapping = rule
            idx = [c["h"] for c in cols].index(header) + 1
            cl = get_column_letter(idx)
            cf_map(ws, f"{cl}{first_row}:{cl}{last_row}", cl, mapping)
        elif kind == "f":
            _, header_or_range, formula, bg, tx = rule[0], rule[1], rule[2], rule[3], rule[4]
            if ":" in header_or_range and header_or_range not in L:
                rng = header_or_range
            else:
                idx = [c["h"] for c in cols].index(header_or_range) + 1
                cl = get_column_letter(idx)
                rng = f"{cl}{first_row}:{cl}{last_row}"
            cf_formula(ws, rng, formula, bg, tx, bold=True)
        elif kind == "bar":
            _, header = rule
            idx = [c["h"] for c in cols].index(header) + 1
            cl = get_column_letter(idx)
            cf_databar(ws, f"{cl}{first_row}:{cl}{last_row}")

    # generic data-quality rules (duplicate id / missing owner / missing due date)
    dq = spec.get("dq", {})
    if dq.get("dup"):
        cl = L[dq["dup"]]
        cf_formula(ws, f"{cl}{first_row}:{cl}{last_row}",
                   f'AND(${cl}{first_row}<>"",COUNTIF(${cl}${first_row}:${cl}${last_row},${cl}{first_row})>1)',
                   R_BG, R_TX, bold=True)
    if dq.get("miss_owner"):
        cl = L[dq["miss_owner"]]
        cf_formula(ws, f"{cl}{first_row}:{cl}{last_row}",
                   f'AND(${L.get(dq.get("key"), cl)}{first_row}<>"",{cl}{first_row}="")',
                   A_BG, A_TX, bold=True)
    if dq.get("miss_due"):
        cl = L[dq["miss_due"]]
        key = dq.get("key")
        cf_formula(ws, f"{cl}{first_row}:{cl}{last_row}",
                   f'AND(${L[key]}{first_row}<>"",{cl}{first_row}="")',
                   A_BG, A_TX, bold=True)

    # freeze + print
    ws.freeze_panes = "A5"
    setup_print(ws, ncols, last_row + spec.get("footer_rows", 0))
    return ws
