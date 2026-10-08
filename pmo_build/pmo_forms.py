"""00_HOME, 03_PROJECT_CHARTER, 15_WEEKLY_STATUS, 21_PMOSOP."""

from openpyxl.styles import Border, Side, Font, Alignment
from openpyxl.utils import get_column_letter

from pmo_core import (NAVY, BLUE, BLUE_MID, BLUE_LT, PALE, WHITE, GRAY_TX, BORDER_C,
                      G_BG, G_TX, A_BG, A_TX, R_BG, R_TX, N_BG, N_TX,
                      font, solid, al, border_all, band, sheet_title, add_list_validation,
                      cf_map, setup_print, write_header, FMT, STATUS_MAP, RAG_MAP)
from pmo_dash_common import (R, PRJ, MS, RAID, ACT, CHG, DEC, QUAL, WBS,
                             PRJ_NAME, PRJ_STATUS, PRJ_RAG, PRJ_PCT, PRJ_APPROVED,
                             PRJ_ACTUAL, PRJ_FORECAST, PRJ_SV, PRJ_MGR, PRJ_HEALTH,
                             PRJ_START, PRJ_BASE, PRJ_FCEND,
                             RAID_TYPE, RAID_STATUS, RAID_DESC, RAID_PRIORITY, RAID_PROJKEY,
                             ACT_DOD, ACT_PROJ, ACT_STATUS,
                             MS_PROJKEY, MS_NAME, MS_FC, MS_STATUS, MS_PROJ,
                             DEC_STATUS, CHG_STATUS,
                             QL_PROJ, QL_RESULT, QL_DEFECTS, QL_STATUS, QL_CORRECTIVE,
                             WBS_PROJ, WBS_STATUS,
                             TICK, merge_style, thin_border, hair_bottom)

SHEETS = [
    ("01", "EXECUTIVE DASHBOARD", "01_EXECUTIVE_DASHBOARD"),
    ("01B", "PROJECT DASHBOARD", "01B_PROJECT_DASHBOARD"),
    ("02", "PROJECT PORTFOLIO", "02_PROJECT_PORTFOLIO"),
    ("03", "PROJECT CHARTER", "03_PROJECT_CHARTER"),
    ("04", "WBS & GANTT", "04_WBS_SCHEDULE"),
    ("05", "MILESTONES", "05_MILESTONES"),
    ("06", "RAID LOG", "06_RAID_LOG"),
    ("07", "ACTION TRACKER", "07_ACTION_TRACKER"),
    ("08", "RACI MATRIX", "08_RACI_MATRIX"),
    ("09", "STAKEHOLDERS", "09_STAKEHOLDERS"),
    ("10", "CHANGE CONTROL", "10_CHANGE_CONTROL"),
    ("11", "COST MANAGEMENT", "11_COST_MANAGEMENT"),
    ("12", "RESOURCE MANAGEMENT", "12_RESOURCE_MANAGEMENT"),
    ("13", "PROCUREMENT", "13_PROCUREMENT"),
    ("14", "QUALITY CONTROL", "14_QUALITY_CONTROL"),
    ("15", "WEEKLY STATUS", "15_WEEKLY_STATUS"),
    ("16", "MEETING MINUTES", "16_MEETING_MINUTES"),
    ("17", "DECISION LOG", "17_DECISION_LOG"),
    ("18", "COMMUNICATION PLAN", "18_COMMUNICATION_PLAN"),
    ("19", "PROJECT CLOSURE", "19_PROJECT_CLOSURE"),
    ("20", "LESSONS LEARNED", "20_LESSONS_LEARNED"),
    ("21", "PMO SOP", "21_PMOSOP"),
    ("22", "30 / 60 / 90 PLAN", "22_30_60_90_PLAN"),
    ("23", "SETTINGS", "23_SETTINGS"),
    ("24", "DATA DICTIONARY", "24_DATA_DICTIONARY"),
]

ORG_GROUPS = [
    ("GOVERNANCE", "03 Project Charter  |  08 RACI Matrix  |  09 Stakeholders  |  "
                   "10 Change Control  |  17 Decision Log  |  18 Communication Plan"),
    ("DELIVERY", "02 Project Portfolio  |  04 WBS & Gantt  |  05 Milestones  |  07 Action Tracker  |  "
                 "12 Resources  |  13 Procurement  |  14 Quality  |  22 30-60-90 Plan"),
    ("CONTROL", "06 RAID Log  |  11 Cost Management  |  15 Weekly Status  |  16 Meeting Minutes  |  "
                "19 Project Closure  |  20 Lessons Learned  |  21 PMO SOP"),
    ("REPORTING", "01 Executive Dashboard  |  01B Project Dashboard  |  23 Settings  |  "
                  "24 Data Dictionary"),
]


# --------------------------------------------------------------------------- #
def build_home(wb):
    ws = wb.create_sheet("00_HOME")
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    NC = 12

    for i in range(1, NC + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14

    sheet_title(
        ws, "ASSISTANT PMO - MANAGEMENT SYSTEM",
        "Control centre for the portfolio: click any tile to open a sheet, then use the "
        "dashboards to brief management.",
        NC,
        note="Nothing here needs typing except the registers themselves - every tile, KPI, "
             "chart and count in this workbook is a formula or a hyperlink.")

    headline = (
        f'="PMO CONTROL CENTRE   |   "&TEXT(TODAY(),"dddd, dd mmmm yyyy")&"   |   "'
        f'&COUNTIF({PRJ_NAME},"?*")&" PROJECTS   |   "&'
        f'COUNTIFS({RAID_TYPE},"Risk",{RAID_STATUS},"<>Closed")&" OPEN RISKS   |   "'
        f'&COUNTIF({ACT_DOD},">0")&" OVERDUE ACTIONS   |   "&COUNTIF({DEC_STATUS},"Pending")'
        f'&" DECISIONS PENDING"'
    )
    merge_style(ws, 4, 1, 4, NC, headline, solid(NAVY), font(11, True, WHITE),
                al("center", "center"))
    ws.row_dimensions[4].height = 26
    ws.row_dimensions[5].height = 8

    band(ws, 6, NC, "NAVIGATION - CLICK ANY TILE TO OPEN THAT SHEET", BLUE)

    tile_border = Border(*[Side(style="thin", color=BLUE_MID)] * 4)
    for i, (num, label, target) in enumerate(SHEETS):
        r = 7 + i // 5
        c1 = 1 + (i % 5) * 2
        cell = merge_style(
            ws, r, c1, r, c1 + 1,
            f'=HYPERLINK("#\'{target}\'!A1","{num}  {label}")',
            solid(BLUE_LT), None, al("center", "center", wrap=True), border=tile_border)
        cell.font = Font(name="Calibri", size=9, bold=True, color="0563C1", underline="single")
    for r in range(7, 12):
        ws.row_dimensions[r].height = 30

    qh = merge_style(ws, 7, 11, 7, 12, "QUICK STATUS", solid(NAVY), font(9, True, WHITE),
                     al("center", "center"), border=thin_border())
    quick = [
        f'=COUNTIF({PRJ_NAME},"?*")&" projects in portfolio"',
        f'=COUNTIFS({RAID_TYPE},"Risk",{RAID_STATUS},"<>Closed")&" open risks"',
        f'=COUNTIF({ACT_DOD},">0")&" overdue actions"',
        f'=COUNTIF({DEC_STATUS},"Pending")&" decisions pending"',
    ]
    for i, formula in enumerate(quick):
        merge_style(ws, 8 + i, 11, 8 + i, 12, formula, solid(PALE), font(9, True, NAVY),
                    al("center", "center", wrap=True), border=thin_border())
    _ = qh

    ws.row_dimensions[12].height = 8
    band(ws, 13, NC, "HOW THIS WORKBOOK IS ORGANISED", BLUE)
    for i, (group, text) in enumerate(ORG_GROUPS):
        r = 14 + i
        cell = merge_style(ws, r, 1, r, NC, f"{group:<12}   {text}",
                           solid(PALE if i % 2 else WHITE), font(9.5),
                           al("left", "center", indent=1), border=hair_bottom())
        ws.row_dimensions[r].height = 19
        cell.alignment = al("left", "center", indent=1)

    ws.row_dimensions[18].height = 8
    band(ws, 19, NC, "TODAY'S PRIORITIES - READ THIS BEFORE THE MANAGEMENT MEETING", BLUE)
    priorities = [
        f'="1.   "&COUNTIF({ACT_DOD},">0")&" actions overdue (worst "'
        f'&MAX({R(ACT, "Days Overdue")})&" days)   -   open 07_ACTION_TRACKER"',
        f'="2.   "&COUNTIFS({MS_FC},">0",{MS_FC},"<"&TODAY(),{MS_STATUS},"<>Complete")'
        f'&" milestones overdue - open 05_MILESTONES"',
        f'="3.   "&COUNTIFS({RAID_TYPE},"Risk",{RAID_PRIORITY},"Critical",{RAID_STATUS},"<>Closed")'
        f'&" critical risks open - open 06_RAID_LOG"',
        f'="4.   "&COUNTIF({CHG_STATUS},"Pending")&" change requests awaiting approval '
        f'- open 10_CHANGE_CONTROL"',
        f'="5.   "&COUNTIF({DEC_STATUS},"Pending")&" decisions pending - open 17_DECISION_LOG"',
    ]
    for i, formula in enumerate(priorities):
        r = 20 + i
        merge_style(ws, r, 1, r, NC, formula,
                    solid(PALE if i % 2 else WHITE), font(10, True, NAVY),
                    al("left", "center", indent=1), border=hair_bottom())
        ws.row_dimensions[r].height = 21

    ws.row_dimensions[25].height = 8
    merge_style(
        ws, 26, 1, 26, NC,
        "EVERYTHING IS AUTOMATIC:  drop-downs come from 23_SETTINGS, KPIs and charts come from the "
        "registers, and the dashboards refresh the moment a register changes.  Sample records are "
        "marked SAMPLE in their Notes column - delete them whenever you like.",
        solid(BLUE_LT), font(9, True, BLUE, italic=True), al("left", "center", indent=1))
    ws.row_dimensions[26].height = 30

    setup_print(ws, NC, 26, title_rows="1:3")
    ws.page_setup.orientation = "landscape"
    return ws


# --------------------------------------------------------------------------- #
def _label(ws, row, text, r1=1, r2=1):
    return merge_style(ws, row, r1, row, r2, text, solid("E4E9EF"), font(9, True, NAVY),
                       al("right", "center"), border=hair_bottom())


def _auto(ws, row, c1, c2, formula, fmt=None, height=20):
    return merge_style(ws, row, c1, row, c2, formula, solid(PALE), font(10, True, NAVY),
                       al("left", "center", indent=1), fmt, border=hair_bottom())


def _input(ws, row, c1, c2, height=40, value=None, fmt=None):
    cell = merge_style(ws, row, c1, row, c2, value, solid(WHITE), font(10),
                       al("left", "top", wrap=True, indent=1), fmt,
                       border=border_all())
    ws.row_dimensions[row].height = height
    return cell


def build_charter(wb):
    ws = wb.create_sheet("03_PROJECT_CHARTER")
    ws.sheet_properties.tabColor = "2E75B6"
    ws.sheet_view.showGridLines = False
    NC = 4
    for col, w in ((1, 30), (2, 48), (3, 34), (4, 28)):
        ws.column_dimensions[get_column_letter(col)].width = w

    sheet_title(
        ws, "PROJECT CHARTER",
        "One-page authorisation document for a single project. Choose the project and the "
        "identification block fills itself; the business case, scope and approvals are yours to write.",
        NC,
        note="Grey cells are automatic (from 02_PROJECT_PORTFOLIO and 06_RAID_LOG). "
             "White cells are for you. Print landscape, one page.")

    band(ws, 4, NC, "SELECT A PROJECT - THE IDENTIFICATION BLOCK BELOW FILLS ITSELF", BLUE)
    _label(ws, 5, "PROJECT")
    ws.merge_cells("B5:D5")
    sel = ws.cell(row=5, column=2, value="ERP Implementation")
    for cc in range(2, 5):
        cell = ws.cell(row=5, column=cc)
        cell.fill = solid(WHITE)
        cell.border = Border(*[Side(style="medium", color=BLUE_MID)] * 4)
    sel.font = font(12, True, NAVY)
    sel.alignment = al("left", "center", indent=1)
    add_list_validation(ws, "B", "L_Projects", 5, 5, first_col_letter="D",
                        prompt="Choose the project this charter belongs to")
    ws.row_dimensions[5].height = 24

    merge_style(ws, 6, 1, 6, NC,
                "LEGEND:   grey cells = automatic from the registers   |   white cells = type here",
                solid(PALE), font(8.5, False, GRAY_TX, italic=True),
                al("left", "center", indent=1))
    ws.row_dimensions[6].height = 16

    m = f'MATCH($B$5,{PRJ_NAME},0)'

    def idx(rng):
        return f'=IFERROR(INDEX({rng},{m}),"-")'

    band(ws, 7, NC, "1.  PROJECT IDENTIFICATION  (AUTOMATIC)", BLUE)
    ident = [
        ("Project ID", idx(R(PRJ, "Project ID")), None),
        ("Customer", idx(R(PRJ, "Customer")), None),
        ("Project Manager", idx(PRJ_MGR), None),
        ("Department", idx(R(PRJ, "Department")), None),
        ("Project Type", idx(R(PRJ, "Project Type")), None),
        ("Priority", idx(R(PRJ, "Priority")), None),
        ("Start Date", idx(PRJ_START), FMT["date"]),
        ("Baseline End Date", idx(PRJ_BASE), FMT["date"]),
        ("Forecast End Date", idx(PRJ_FCEND), FMT["date"]),
        ("Total Approved Amount", idx(PRJ_APPROVED), "#,##0"),
        ("Actual Cost", idx(PRJ_ACTUAL), "#,##0"),
        ("Status", idx(PRJ_STATUS), None),
        ("RAG", idx(PRJ_RAG), None),
    ]
    for i, (label, formula, fmt) in enumerate(ident):
        r = 8 + i
        _label(ws, r, label)
        _auto(ws, r, 2, 4, formula, fmt, height=19)
    cf_map(ws, "B20:B20", "B", RAG_MAP, first_row=20)
    cf_map(ws, "B19:B19", "B", STATUS_MAP, first_row=19)

    band(ws, 21, NC, "2.  BUSINESS CASE  (TYPE HERE)", BLUE)
    for i, label in enumerate(["Problem / Opportunity", "Business Objective",
                               "Expected Benefits", "Success Criteria"]):
        r = 22 + i
        _label(ws, r, label)
        _input(ws, r, 2, 4, height=38)

    band(ws, 26, NC, "3.  SCOPE  (TYPE HERE)", BLUE)
    for i, label in enumerate(["In Scope", "Out of Scope", "Key Deliverables"]):
        r = 27 + i
        _label(ws, r, label)
        _input(ws, r, 2, 4, height=38)

    band(ws, 30, NC, "4.  APPROACH & GOVERNANCE", BLUE)
    _label(ws, 31, "Delivery Approach")
    _input(ws, 31, 2, 4, height=38)
    _label(ws, 32, "Key Assumptions")
    _input(ws, 32, 2, 4, height=38)
    _label(ws, 33, "Top Open Risk (automatic)")
    _auto(ws, 33, 2, 4,
          f'=IFERROR(INDEX({RAID_DESC},MATCH("Risk|"&$B$5&"|1",{RAID_PROJKEY},0)),'
          f'"No open risk recorded")', None, height=34)
    _label(ws, 34, "Reporting & Governance")
    _input(ws, 34, 2, 4, height=38)

    band(ws, 35, NC, "5.  APPROVALS  (TYPE HERE)", BLUE)
    for i, label in enumerate(["Prepared By", "Reviewed By", "Approved By",
                               "Approval Date", "Signature"]):
        r = 36 + i
        _label(ws, r, label)
        _input(ws, r, 2, 4, height=24,
               fmt=FMT["date"] if label == "Approval Date" else None)

    setup_print(ws, NC, 40, title_rows="1:3")
    ws.page_setup.orientation = "landscape"
    return ws


# --------------------------------------------------------------------------- #
def build_weekly_status(wb):
    ws = wb.create_sheet("15_WEEKLY_STATUS")
    ws.sheet_properties.tabColor = "1F4E79"
    ws.sheet_view.showGridLines = False
    NC = 8
    for col, w in ((1, 22), (2, 26), (3, 26), (4, 24), (5, 18), (6, 18), (7, 18), (8, 16)):
        ws.column_dimensions[get_column_letter(col)].width = w

    sheet_title(
        ws, "WEEKLY PROJECT STATUS REPORT",
        "One-page, print-ready status report. Choose a project: the schedule, cost, quality, "
        "risk, issue and action sections fill themselves; the text boxes are yours.",
        NC,
        note="Print landscape on one page. The automatic sections read from the registers, "
             "so the report can never disagree with the dashboard.")

    P = "$E$5"
    m = f'MATCH({P},{PRJ_NAME},0)'

    def idx(rng, fallback='-'):
        return f'IFERROR(INDEX({rng},{m}),"{fallback}")'

    band(ws, 4, NC, "REPORT HEADER", BLUE)

    def hdr_pair(row, l1, l2, label, v1, v2, value, fmt=None, fill=WHITE):
        merge_style(ws, row, l1, row, l2, label, solid("E4E9EF"), font(9, True, NAVY),
                    al("right", "center"), border=hair_bottom())
        cell = merge_style(ws, row, v1, row, v2, value, solid(fill), font(10, True, NAVY),
                           al("left", "center", indent=1), fmt, border=hair_bottom())
        return cell

    hdr_pair(5, 1, 1, "REPORTING WEEK",
             2, 3, '="Week ending "&TEXT(TODAY()-WEEKDAY(TODAY(),2)+7,"dd mmm yyyy")')
    hdr_pair(5, 4, 4, "PROJECT", 5, 6, "ERP Implementation")
    add_list_validation(ws, "E", "L_Projects", 5, 5, first_col_letter="F",
                        prompt="Choose the project this report covers")
    hdr_pair(5, 7, 7, "REPORT DATE", 8, 8, "=TODAY()", FMT["date"])
    ws.row_dimensions[5].height = 22

    hdr_pair(6, 1, 1, "PROJECT MANAGER", 2, 3, f'={idx(PRJ_MGR)}')
    hdr_pair(6, 4, 4, "STATUS", 5, 6, f'={idx(PRJ_STATUS)}')
    hdr_pair(6, 7, 7, "RAG", 8, 8, f'={idx(PRJ_RAG)}')
    ws.row_dimensions[6].height = 22

    hdr_pair(7, 1, 1, "CUSTOMER", 2, 3, f'={idx(R(PRJ, "Customer"))}')
    hdr_pair(7, 4, 4, "% COMPLETE", 5, 6, f'={idx(PRJ_PCT)}', FMT["pct"])
    hdr_pair(7, 7, 7, "HEALTH", 8, 8, f'={idx(PRJ_HEALTH)}')
    ws.row_dimensions[7].height = 22
    cf_map(ws, "E6:E6", "E", STATUS_MAP, first_row=6)
    cf_map(ws, "H6:H6", "H", RAG_MAP, first_row=6)
    ws.row_dimensions[8].height = 6

    def section(row, title):
        band(ws, row, NC, title, BLUE_MID, size=9.5)

    def note(row, height):
        ws.row_dimensions[row].height = height

    def auto(row, formula, height=20):
        merge_style(ws, row, 1, row, NC, formula, solid(PALE), font(9.5, True, NAVY),
                    al("left", "center", indent=1), border=hair_bottom())
        ws.row_dimensions[row].height = height

    def box(row, height):
        merge_style(ws, row, 1, row, NC, None, solid(WHITE), font(10),
                    al("left", "top", wrap=True, indent=1), border=border_all())
        ws.row_dimensions[row].height = height

    r = 9
    section(r, "EXECUTIVE SUMMARY   -   TYPE HERE")
    box(r + 1, 54); r += 2
    section(r, "COMPLETED THIS WEEK   -   TYPE HERE")
    box(r + 1, 44); r += 2
    section(r, "PLANNED NEXT WEEK   -   TYPE HERE")
    box(r + 1, 44); r += 2

    section(r, "KEY MILESTONES   -   AUTOMATIC (NEXT 3 FOR THIS PROJECT)")
    for k in range(1, 4):
        auto(r + k,
             f'="{TICK}   "&IFERROR(INDEX({MS_NAME},MATCH({P}&"|{k}",{MS_PROJKEY},0))&'
             f'"   -   forecast "&TEXT(INDEX({MS_FC},MATCH({P}&"|{k}",{MS_PROJKEY},0)),"dd mmm yyyy")&'
             f'"   -   "&INDEX({MS_STATUS},MATCH({P}&"|{k}",{MS_PROJKEY},0)),'
             f'"No further milestone recorded")', height=18)
    r += 4

    section(r, "SCHEDULE STATUS   -   AUTOMATIC")
    auto(r + 1,
         f'="Schedule variance: "&IFERROR(INDEX({PRJ_SV},{m}),"-")&" days   |   Tasks complete: "&'
         f'IF(COUNTIF({WBS_PROJ},{P})=0,"-",COUNTIFS({WBS_PROJ},{P},{WBS_STATUS},"Complete")&'
         f'" of "&COUNTIF({WBS_PROJ},{P}))&"   |   Milestones overdue: "&'
         f'COUNTIFS({MS_PROJ},{P},{MS_FC},">0",{MS_FC},"<"&TODAY(),{MS_STATUS},"<>Complete")&'
         f'"   |   Milestones due in 7 days: "&'
         f'COUNTIFS({MS_PROJ},{P},{MS_FC},">="&TODAY(),{MS_FC},"<="&TODAY()+7,{MS_STATUS},"<>Complete")')
    r += 2

    section(r, "COST STATUS   -   AUTOMATIC")
    auto(r + 1,
         f'="Approved: "&TEXT(IFERROR(INDEX({PRJ_APPROVED},{m}),0),"#,##0")&"   |   Actual: "&'
         f'TEXT(IFERROR(INDEX({PRJ_ACTUAL},{m}),0),"#,##0")&" ("&'
         f'TEXT(IFERROR(INDEX({PRJ_ACTUAL},{m})/INDEX({PRJ_APPROVED},{m}),0),"0%")&")   |   Forecast: "&'
         f'TEXT(IFERROR(INDEX({PRJ_FORECAST},{m}),0),"#,##0")&IF('
         f'IFERROR(INDEX({PRJ_FORECAST},{m}),0)>IFERROR(INDEX({PRJ_APPROVED},{m}),0),'
         f'"   -   OVER APPROVED BY "&TEXT(INDEX({PRJ_FORECAST},{m})-INDEX({PRJ_APPROVED},{m}),"#,##0"),'
         f'"   -   within approved")')
    r += 2

    section(r, "SCOPE STATUS   -   TYPE HERE")
    box(r + 1, 44); r += 2

    section(r, "QUALITY STATUS   -   AUTOMATIC")
    auto(r + 1,
         f'="Checks passed: "&COUNTIFS({QL_PROJ},{P},{QL_RESULT},"Pass")&" of "&'
         f'COUNTIF({QL_PROJ},{P})&"   |   Failed: "&COUNTIFS({QL_PROJ},{P},{QL_RESULT},"Fail")&'
         f'"   |   Defects: "&SUMIFS({QL_DEFECTS},{QL_PROJ},{P})&"   |   Open corrective actions: "&'
         f'COUNTIFS({QL_PROJ},{P},{QL_CORRECTIVE},"<>",{QL_STATUS},"<>Complete")&'
         f'"   |   Pending decisions: "&COUNTIFS({R(DEC, "Project")},{P},{DEC_STATUS},"Pending")')
    r += 2

    section(r, "TOP RISKS   -   AUTOMATIC (TOP 3 BY SCORE)")
    for k in range(1, 4):
        auto(r + k,
             f'="{TICK}   "&IFERROR(INDEX({RAID_DESC},MATCH("Risk|"&{P}&"|{k}",{RAID_PROJKEY},0))&'
             f'"   ["&INDEX({RAID_PRIORITY},MATCH("Risk|"&{P}&"|{k}",{RAID_PROJKEY},0))&"]   owner: "&'
             f'INDEX({R(RAID, "Owner")},MATCH("Risk|"&{P}&"|{k}",{RAID_PROJKEY},0)),"")', height=18)
    r += 4

    section(r, "TOP ISSUES   -   AUTOMATIC (TOP 2)")
    for k in range(1, 3):
        auto(r + k,
             f'="{TICK}   "&IFERROR(INDEX({RAID_DESC},MATCH("Issue|"&{P}&"|{k}",{RAID_PROJKEY},0))&'
             f'"   owner: "&INDEX({R(RAID, "Owner")},MATCH("Issue|"&{P}&"|{k}",{RAID_PROJKEY},0)),"")',
             height=18)
    r += 3

    section(r, "DECISIONS REQUIRED FROM MANAGEMENT   -   AUTOMATIC COUNTS + YOUR ASK")
    auto(r + 1,
         f'="Pending decisions: "&COUNTIFS({R(DEC, "Project")},{P},{DEC_STATUS},"Pending")&'
         f'"   |   Pending changes: "&COUNTIFS({R(CHG, "Project")},{P},{CHG_STATUS},"Pending")&'
         f'"   |   Overdue actions: "&COUNTIFS({ACT_PROJ},{P},{ACT_DOD},">0")&'
         f'"   |   Open risks: "&COUNTIFS({R(RAID, "Project")},{P},{RAID_TYPE},"Risk",{RAID_STATUS},"<>Closed")')
    box(r + 2, 40)
    r += 3

    section(r, "SUPPORT REQUIRED   -   TYPE HERE")
    box(r + 1, 40)
    r += 2

    setup_print(ws, NC, r, title_rows="1:3")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToHeight = 1
    return ws


# --------------------------------------------------------------------------- #
_SOP_DAILY = [
    ("09:00", "Portfolio sweep", "Open the executive dashboard and read the headline strip, "
                                 "overdue actions and red projects", "List of today's chases",
     "Assistant PMO"),
    ("09:15", "Team stand-up", "Capture every blocker as an action with one owner and one due date",
     "Actions logged in 07_ACTION_TRACKER", "Project Manager"),
    ("10:00", "Action chasing", "Chase every action due today or overdue; update status in the tracker",
     "Action statuses refreshed", "Assistant PMO"),
    ("11:00", "RAID review", "Update open risks and issues, re-score anything that changed",
     "RAID log current", "Assistant PMO"),
    ("14:00", "Data quality", "Check duplicate IDs, missing owners and missing dates flagged red or "
                              "amber in the registers", "Clean registers", "Assistant PMO"),
    ("16:00", "Change control", "Check the compliance column: no change may be implemented without "
                                "approval", "Compliance column = OK", "Assistant PMO"),
    ("17:00", "Tomorrow's plan", "Note the meetings, decisions and deliveries due tomorrow",
     "Prepared day plan", "Assistant PMO"),
]

_SOP_WEEKLY = [
    ("Monday 10:00", "Project manager call", "Each PM reports status, next milestone, top risk and "
                                             "the help they need", "Status inputs from every PM",
     "Assistant PMO"),
    ("Tuesday", "Register refresh", "Update portfolio, milestones, cost and actions from PM inputs",
     "Registers current for the week", "Assistant PMO"),
    ("Wednesday", "Quality and procurement check", "Review quality results, defects and delivery "
                                                   "of ordered items", "Quality and procurement status",
     "Assistant PMO"),
    ("Thursday", "Draft the status report", "Complete 15_WEEKLY_STATUS for each active project; "
                                            "automatic sections must not be overwritten", "Draft reports",
     "Assistant PMO"),
    ("Friday 09:00", "Publish the dashboard", "Refresh the executive dashboard and confirm every KPI "
                                              "and chart is current", "Published dashboard",
     "Assistant PMO"),
    ("Friday", "Executive pack", "Send the dashboard link with the red projects, overdue actions "
                                 "and decisions required", "Executive communication", "Assistant PMO"),
    ("Friday", "Action close-out", "Close completed actions; escalate anything overdue by more "
                                   "than 7 days", "Updated action tracker", "Assistant PMO"),
    ("Friday", "Decision and change sweep", "List pending decisions and changes older than 5 days "
                                            "for the steering meeting", "Decision agenda", "Assistant PMO"),
]

_SOP_MONTHLY = [
    ("Working day 1", "Portfolio review", "Review every project's RAG, schedule variance and cost "
                                          "variance with the PMs", "Portfolio status agreed",
     "PMO Lead"),
    ("Working day 2", "Cost report", "Reconcile approved amount, actual and forecast with Finance",
     "Cost report issued", "Assistant PMO"),
    ("Working day 3", "Risk deep dive", "Score every open risk, confirm mitigations and escalations",
     "Risk report issued", "PMO Lead"),
    ("First Thursday", "Steering committee", "Present the dashboard: red projects, decisions "
                                             "required, change requests and benefits",
     "Steering minutes and decisions", "PMO Lead"),
    ("Last Friday", "Lessons learned", "Capture what went well and what did not; assign the "
                                       "improvement actions", "Updated lessons learned",
     "Assistant PMO"),
    ("Last Friday", "30/60/90 review", "Update 22_30_60_90_PLAN progress and next month's targets",
     "Plan status updated", "Assistant PMO"),
    ("Month end", "Archive and hygiene", "Archive completed projects, purge stale sample data, "
                                         "confirm drop-downs still match 23_SETTINGS",
     "Clean workbook", "Assistant PMO"),
]

_SOP_COLS = [
    {"h": "WHEN", "w": 16}, {"h": "ACTIVITY", "w": 34},
    {"h": "WHAT GOOD LOOKS LIKE", "w": 56}, {"h": "OUTPUT", "w": 34},
    {"h": "OWNER", "w": 18},
]


def build_sop(wb):
    ws = wb.create_sheet("21_PMOSOP")
    ws.sheet_properties.tabColor = "2E75B6"
    ws.sheet_view.showGridLines = False
    NC = 5
    for i, c in enumerate(_SOP_COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = c["w"]

    sheet_title(
        ws, "PMO STANDARD OPERATING PROCEDURES",
        "The cadence that keeps the system honest: daily hygiene, the weekly reporting cycle and "
        "the monthly governance rhythm.",
        NC,
        note="Print this sheet and keep it next to the workbook. If the routine stops, the "
             "dashboard stops being true.")

    ws.row_dimensions[4].height = 8
    row = 5
    for title, rows in (("DAILY ROUTINE", _SOP_DAILY),
                        ("WEEKLY ROUTINE", _SOP_WEEKLY),
                        ("MONTHLY ROUTINE", _SOP_MONTHLY)):
        band(ws, row, NC, title, BLUE)
        row += 1
        write_header(ws, _SOP_COLS, row=row, height=26)
        row += 1
        first = row
        for data in rows:
            for i, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=i, value=value)
                cell.font = font(9.5, i <= 2)
                cell.alignment = al("left", "top", wrap=True, indent=1 if i > 1 else 0)
            ws.row_dimensions[row].height = 28
            row += 1
        from openpyxl.worksheet.table import Table, TableStyleInfo
        tab = Table(displayName=f"tblSOP{title.split()[0].title()}",
                     ref=f"A{first - 1}:{get_column_letter(NC)}{row - 1}")
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False,
                                            showLastColumn=False, showRowStripes=True,
                                            showColumnStripes=False)
        ws.add_table(tab)
        ws.row_dimensions[row].height = 8
        row += 1

    setup_print(ws, NC, row, title_rows="1:3")
    ws.page_setup.orientation = "landscape"
    return ws
