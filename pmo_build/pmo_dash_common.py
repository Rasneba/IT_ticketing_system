"""Shared range handles + building blocks for the PMO dashboards."""

from openpyxl.styles import Border, Side, Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

from pmo_core import (NAVY, BLUE, BLUE_MID, BLUE_LT, PALE, WHITE, GRAY_TX, BORDER_C,
                      G_BG, G_TX, A_BG, A_TX, R_BG, R_TX,
                      font, solid, al, band, sheet_title, add_list_validation,
                      cf_map, cf_databar, setup_print, col_letter_map, FMT,
                      STATUS_MAP, RAG_MAP, unmerge_overlap)
from pmo_specs import PORTFOLIO_SPEC, WBS_SPEC, MILESTONE_SPEC, RAID_SPEC, ACTION_SPEC
from pmo_specs2 import (CHANGE_SPEC, DECISION_SPEC, COST_SPEC, QUALITY_SPEC, CLOSURE_SPEC)

_MAPS = {}


def colmap(spec):
    if spec["sheet"] not in _MAPS:
        _MAPS[spec["sheet"]] = col_letter_map(spec["cols"])
    return _MAPS[spec["sheet"]]


def R(spec, header, r1=5, r2=34):
    letter = colmap(spec)[header]
    return f"'{spec['sheet']}'!${letter}${r1}:${letter}${r2}"


PRJ = PORTFOLIO_SPEC
WBS = WBS_SPEC
MS = MILESTONE_SPEC
RAID = RAID_SPEC
ACT = ACTION_SPEC
CHG = CHANGE_SPEC
DEC = DECISION_SPEC
COST = COST_SPEC
QUAL = QUALITY_SPEC
CLO = CLOSURE_SPEC

PRJ_ID = R(PRJ, "Project ID")
PRJ_NAME = R(PRJ, "Project Name")
PRJ_STATUS = R(PRJ, "Status")
PRJ_RAG = R(PRJ, "RAG")
PRJ_MGR = R(PRJ, "Project Manager")
PRJ_DEPT = R(PRJ, "Department")
PRJ_START = R(PRJ, "Start Date")
PRJ_BASE = R(PRJ, "Baseline End Date")
PRJ_FCEND = R(PRJ, "Forecast End Date")
PRJ_PCT = R(PRJ, "% Complete")
PRJ_APPROVED = R(PRJ, "Total Approved Amount")
PRJ_ACTUAL = R(PRJ, "Actual Cost")
PRJ_FORECAST = R(PRJ, "Forecast Cost")
PRJ_SV = R(PRJ, "Schedule Variance")
PRJ_CV = R(PRJ, "Cost Variance")
PRJ_NEXTMS = R(PRJ, "Next Milestone")
PRJ_MSDATE = R(PRJ, "Milestone Date")
PRJ_HEALTH = R(PRJ, "Project Health")
PRJ_TIMELINE = R(PRJ, "Timeline")
PRJ_BA = R(PRJ, "Budget/Actual")
PRJ_CHARTIDX = R(PRJ, "ChartIdx")
PRJ_REDRANK = R(PRJ, "RedRank")
PRJ_VIEWRANK = R(PRJ, "ViewRank")

RAID_TYPE = R(RAID, "Type")
RAID_STATUS = R(RAID, "Status")
RAID_PRIORITY = R(RAID, "Priority")
RAID_DESC = R(RAID, "Description")
RAID_PROJ = R(RAID, "Project")
RAID_RANK = R(RAID, "Rank")
RAID_PROJKEY = R(RAID, "ProjKey")
RAID_ESC = R(RAID, "Escalation")

ACT_STATUS = R(ACT, "Status")
ACT_DOD = R(ACT, "Days Overdue")
ACT_RANK = R(ACT, "OverdueRank")
ACT_PROJKEY = R(ACT, "ProjKey")
ACT_OWNER = R(ACT, "Owner")
ACT_ACTION = R(ACT, "Action")
ACT_DUE = R(ACT, "Due Date")
ACT_PROJ = R(ACT, "Project")

MS_KEY = R(MS, "Key")
MS_NAME = R(MS, "Milestone")
MS_PROJ = R(MS, "Project")
MS_OWNER = R(MS, "Owner")
MS_FC = R(MS, "Forecast Date")
MS_BASE = R(MS, "Baseline Date")
MS_ACTUAL = R(MS, "Actual Date")
MS_STATUS = R(MS, "Status")
MS_RAG = R(MS, "RAG")
MS_DAYS = R(MS, "Days Remaining")
MS_PROJKEY = R(MS, "ProjKey")

CHG_STATUS = R(CHG, "Approval Status")
CHG_RANK = R(CHG, "PendingRank")
CHG_ID = R(CHG, "Change ID")
CHG_DESC = R(CHG, "Change Description")
CHG_COST = R(CHG, "Cost Impact")

DEC_STATUS = R(DEC, "Status")
DEC_DUE = R(DEC, "Due Date")

WBS_PROJ = R(WBS, "Project")
WBS_STATUS = R(WBS, "Status")
WBS_OWNER = R(WBS, "Owner")
WBS_PCT = R(WBS, "% Complete")

CST_PROJ = R(COST, "Project")
CST_CAT = R(COST, "Cost Category")
CST_APPROVED = R(COST, "Total Approved Amount")
CST_ACTUAL = R(COST, "Actual Cost")
CST_FORECAST = R(COST, "Forecast Cost")

QL_PROJ = R(QUAL, "Project")
QL_RESULT = R(QUAL, "Result")
QL_STATUS = R(QUAL, "Status")
QL_DEFECTS = R(QUAL, "Defects")
QL_PLANNED = R(QUAL, "Planned Date")
QL_ACTUALDATE = R(QUAL, "Actual Date")
QL_ID = R(QUAL, "QC ID")
QL_CORRECTIVE = R(QUAL, "Corrective Action")

CL_ID = R(CLO, "Item ID")
CL_STATUS = R(CLO, "Status")
CL_DUE = R(CLO, "Due Date")

TICK = "\u2022"


def merge_style(ws, r1, c1, r2, c2, value=None, fill=None, fnt=None, alignment=None,
                fmt=None, border=None):
    if (r1, c1) != (r2, c2):
        unmerge_overlap(ws, r1, c1, r2, c2)
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    for rr in range(r1, r2 + 1):
        for cc in range(c1, c2 + 1):
            cell = ws.cell(row=rr, column=cc)
            if fill:
                cell.fill = fill
            if border:
                cell.border = border
    cell = ws.cell(row=r1, column=c1)
    if value is not None:
        cell.value = value
    if fnt:
        cell.font = fnt
    if alignment:
        cell.alignment = alignment
    if fmt:
        cell.number_format = fmt
    return cell


def thin_border(color=BLUE_MID):
    return Border(*[Side(style="thin", color=color)] * 4)


def hair_bottom():
    return Border(bottom=Side(style="hair", color=BORDER_C))


def kpi_card(ws, row, col, label, value, caption, fmt=None, alert=None):
    a = get_column_letter(col)
    bd = thin_border()
    merge_style(ws, row, col, row, col + 1, label, solid(BLUE),
                font(8.5, True, WHITE), al("center", "center", wrap=True), border=bd)
    v = merge_style(ws, row + 1, col, row + 1, col + 1, value, solid(PALE),
                    font(19, True, NAVY), al("center", "center"), fmt or "General", border=bd)
    merge_style(ws, row + 2, col, row + 2, col + 1, caption, solid(PALE),
                font(8, False, GRAY_TX, italic=True), al("center", "center", wrap=True), border=bd)
    ws.row_dimensions[row].height = 17
    ws.row_dimensions[row + 1].height = 30
    ws.row_dimensions[row + 2].height = 15
    if alert:
        addr = f"{a}{row + 1}"
        bg, tx = (R_BG, R_TX) if alert == "red" else (A_BG, A_TX)
        ws.conditional_formatting.add(addr, FormulaRule(
            formula=[f"{a}{row + 1}>0"], fill=solid(bg), font=font(19, True, tx), stopIfTrue=True))
        ws.conditional_formatting.add(addr, FormulaRule(
            formula=[f"{a}{row + 1}=0"], fill=solid(G_BG), font=font(19, True, G_TX), stopIfTrue=True))
    return v


def panel_header(ws, row, c1, c2, text):
    merge_style(ws, row, c1, row, c2, text, solid(BLUE_MID), font(9.5, True, WHITE),
                al("left", "center", indent=1), border=thin_border(BLUE))
    ws.row_dimensions[row].height = 19


def panel_line(ws, row, c1, c2, formula, shade=False, height=16):
    fill = solid(PALE) if shade else solid(WHITE)
    merge_style(ws, row, c1, row, c2, formula, fill, font(9),
                al("left", "center", indent=1), border=hair_bottom())
    ws.row_dimensions[row].height = height


def cd_label(ws, row, text):
    merge_style(ws, row, 16, row, 20, text, solid(BLUE_MID), font(9, True, WHITE),
                al("left", "center", indent=1))


def cd_header(ws, row, labels):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=16 + i, value=h)
        c.font = font(9, True, WHITE)
        c.fill = solid(NAVY)
        c.alignment = al("center", "center", wrap=True)
    ws.row_dimensions[row].height = 26


def cd_value(ws, row, col_off, value, fmt=None):
    c = ws.cell(row=row, column=16 + col_off, value=value)
    c.font = font(9)
    c.alignment = al("left", "center")
    if fmt:
        c.number_format = fmt
    return c


def href(ws, min_col, min_row, max_col, max_row):
    from openpyxl.chart import Reference
    return Reference(ws, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)


def link_cell(ws, row, c1, c2, formula, fill, width_font=9):
    cell = merge_style(ws, row, c1, row, c2, formula, fill, None,
                       al("center", "center"), border=thin_border())
    cell.font = Font(name="Calibri", size=width_font, bold=True, color="0563C1", underline="single")
    return cell
