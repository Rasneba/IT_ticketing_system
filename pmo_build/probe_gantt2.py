"""Progressively strip the gantt sheet to find the offending element."""

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
import win32com.client as win32
import build as B
import pmo_specs as S1

TEMP = os.path.join(os.environ["TEMP"], "pmo_g3.xlsx")
customers = B.load_customers()
N = len(S1.WBS_SPEC["cols"])


def base():
    wb = Workbook()
    wb.remove(wb.active)
    for spec in B.SPECS:
        B.build_table_sheet(wb, spec)
    B.build_settings(wb, customers)
    return wb


def strip_merges(ws):
    for rng in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(rng))


def strip_beyond(ws):
    for r in range(1, 60):
        for c in range(N + 1, N + 26):
            cell = ws.cell(row=r, column=c)
            cell.value = None


def strip_cf(ws):
    ws.conditional_formatting._cf_rules = {}


def strip_widths(ws):
    for c in range(N + 1, N + 26):
        ws.column_dimensions.pop(get_column_letter(c), None)


def strip_print(ws):
    ws.print_area = None
    ws.print_title_rows = None


def strip_rows(ws):
    for r in (3, 4, 36):
        ws.row_dimensions.pop(r, None)


def run(tag, muts):
    wb = base()
    ws = wb["04_WBS_SCHEDULE"]
    B.add_gantt(ws)
    for m in muts:
        m(ws)
    wb.calculation.fullCalcOnLoad = True
    wb.save(TEMP)
    try:
        bk = xl.Workbooks.Open(TEMP)
        print(f"OK   {tag}")
        bk.Close(False)
    except Exception:
        print(f"FAIL {tag}")


xl = win32.Dispatch("Excel.Application")
xl.Visible = False
xl.DisplayAlerts = False
try:
    run("full", [])
    run("strip-beyond", [strip_merges, strip_beyond])
    run("strip-beyond+cf", [strip_merges, strip_beyond, strip_cf])
    run("strip-all", [strip_merges, strip_beyond, strip_cf, strip_widths, strip_print, strip_rows])
finally:
    xl.Quit()
