"""01B_PROJECT_DASHBOARD - single-project view with selector, KPIs, panels and charts."""

from openpyxl.styles import Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

from pmo_core import (NAVY, BLUE, BLUE_MID, PALE, WHITE, GRAY_TX,
                      G_BG, G_TX, A_BG, A_TX, R_BG, R_TX,
                      font, solid, al, band, add_list_validation, cf_map, cf_databar,
                      setup_print, FMT, STATUS_MAP, RAG_MAP)
from pmo_charts import doughnut, bar
from pmo_dash_common import (R, PRJ, MS, RAID, ACT, COST, WBS, QUAL, DEC, CHG,
                             PRJ_NAME, PRJ_STATUS, PRJ_RAG, PRJ_HEALTH, PRJ_START,
                             PRJ_BASE, PRJ_FCEND, PRJ_PCT, PRJ_APPROVED, PRJ_ACTUAL,
                             PRJ_FORECAST, PRJ_SV, PRJ_NEXTMS, PRJ_MSDATE,
                             PRJ_MGR, PRJ_DEPT,
                             RAID_PROJKEY, RAID_DESC, RAID_PRIORITY, RAID_PROJ,
                             RAID_TYPE, RAID_STATUS, RAID_RANK,
                             ACT_PROJKEY, ACT_OWNER, ACT_ACTION, ACT_DUE, ACT_DOD, ACT_STATUS,
                             ACT_PROJ,
                             MS_PROJKEY, MS_NAME, MS_FC, MS_STATUS, MS_PROJ,
                             WBS_PROJ, WBS_STATUS, WBS_OWNER,
                             CST_PROJ, CST_CAT, CST_APPROVED, CST_ACTUAL, CST_FORECAST,
                             DEC_STATUS, CHG_STATUS,
                             TICK, merge_style, kpi_card, panel_header, panel_line,
                             cd_label, cd_header, cd_value, href, thin_border, hair_bottom)

OWNERS = ["Ahmed Khan", "Lisa Chen", "David Okoro", "Maria Silva",
          "Priya Raman", "Tom Becker", "Grace Alemu", "Sarah Mitchell"]
SEL = "$D$3"


def build_project_dashboard(wb):
    ws = wb.create_sheet("01B_PROJECT_DASHBOARD")
    ws.sheet_properties.tabColor = "2E75B6"
    ws.sheet_view.showGridLines = False
    NC = 14

    for i in range(1, NC + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    for i in range(16, 20):
        ws.column_dimensions[get_column_letter(i)].width = 18
    ws.column_dimensions["O"].width = 3

    m = f'MATCH({SEL},{PRJ_NAME},0)'

    def idx(rng):
        return f'IFERROR(INDEX({rng},{m}),"-")'

    # --------------------------------------------------------------- title --- #
    merge_style(ws, 1, 1, 1, NC, "PROJECT DASHBOARD", solid(NAVY), font(16, True, WHITE),
                al("left", "center", indent=1))
    ws.row_dimensions[1].height = 32
    merge_style(ws, 2, 1, 2, NC,
                f'="SINGLE PROJECT VIEW   -   "&IF({SEL}="","SELECT A PROJECT IN ROW 3",{SEL})',
                solid(BLUE), font(9.5, False, WHITE, italic=True), al("left", "center", indent=1))
    ws.row_dimensions[2].height = 17
    merge_style(ws, 3, 1, 3, 3, "SELECT PROJECT:", solid(BLUE_MID), font(9.5, True, WHITE),
                al("right", "center"))
    ws.merge_cells("D3:H3")
    sel = ws.cell(row=3, column=4, value="ERP Implementation")
    for cc in range(4, 9):
        cell = ws.cell(row=3, column=cc)
        cell.fill = solid(WHITE)
        cell.border = Border(*[Side(style="medium", color=BLUE_MID)] * 4)
    sel.font = font(12, True, NAVY)
    sel.alignment = al("left", "center", indent=1)
    add_list_validation(ws, "D", "L_Projects", 3, 3, first_col_letter="H",
                        prompt="Choose the project you want to review")
    merge_style(ws, 3, 9, 3, NC,
                f'="AS AT "&TEXT(TODAY(),"dd mmm yyyy")&"   |   LAST UPDATED: "&'
                f'IFERROR(TEXT(INDEX({R(PRJ, "Last Updated")},{m}),"dd mmm yyyy"),"-")',
                solid(PALE), font(9, True, GRAY_TX), al("center", "center"))
    ws.row_dimensions[3].height = 26

    band(ws, 4, NC, "PROJECT SNAPSHOT - EVERY FIGURE FOLLOWS THE PROJECT SELECTED ABOVE", BLUE)

    # ----------------------------------------------------------- KPI cards --- #
    k1 = [
        (1, "STATUS", f'={idx(PRJ_STATUS)}', "LIFECYCLE STATUS", None, None),
        (3, "RAG", f'={idx(PRJ_RAG)}', "OVERALL HEALTH", None, None),
        (5, "% COMPLETE", f'={idx(PRJ_PCT)}', "REPORTED BY THE PM", FMT["pct"], None),
        (7, "SCHEDULE VARIANCE", f'={idx(PRJ_SV)}', "FORECAST vs BASELINE (DAYS)", "0", None),
        (9, "APPROVED AMOUNT", f'={idx(PRJ_APPROVED)}', "TOTAL APPROVED", "#,##0", None),
        (11, "ACTUAL COST", f'={idx(PRJ_ACTUAL)}', "SPENT TO DATE", "#,##0", None),
        (13, "FORECAST COST", f'={idx(PRJ_FORECAST)}', "EXPECTED AT COMPLETION", "#,##0", None),
    ]
    for col, lab, val, cap, fmt, alert in k1:
        kpi_card(ws, 5, col, lab, val, cap, fmt, alert)

    k2 = [
        (1, "PROJECT HEALTH", f'={idx(PRJ_HEALTH)}', "AUTOMATIC FLAG", None, None),
        (3, "CUSTOMER", f'={idx(R(PRJ, "Customer"))}', "CLIENT / INTERNAL", None, None),
        (5, "PROJECT MANAGER", f'={idx(PRJ_MGR)}', "ACCOUNTABLE OWNER", None, None),
        (7, "DEPARTMENT", f'={idx(PRJ_DEPT)}', "OWNING DEPARTMENT", None, None),
        (9, "PRIORITY", f'={idx(R(PRJ, "Priority"))}', "BUSINESS PRIORITY", None, None),
        (11, "START DATE", f'={idx(PRJ_START)}', "PLANNED START", FMT["date"], None),
        (13, "BASELINE END", f'={idx(PRJ_BASE)}', "CONTRACTUAL FINISH", FMT["date"], None),
    ]
    for col, lab, val, cap, fmt, alert in k2:
        kpi_card(ws, 9, col, lab, val, cap, fmt, alert)

    cf_map(ws, "A6:A6", "A", STATUS_MAP, first_row=6)
    cf_map(ws, "C6:C6", "C", RAG_MAP, first_row=6)
    cf_map(ws, "A10:A10", "A",
           {"Healthy": (G_BG, G_TX), "Needs Attention": (A_BG, A_TX), "Critical": (R_BG, R_TX),
            "Closed": ("E4E9EF", "3B4A59"), "N/A": ("E4E9EF", "3B4A59")}, first_row=10)
    ws.conditional_formatting.add("G6", FormulaRule(
        formula=['AND(ISNUMBER(G6),G6>7)'], fill=solid(R_BG), font=font(19, True, R_TX),
        stopIfTrue=True))
    ws.conditional_formatting.add("G6", FormulaRule(
        formula=['AND(ISNUMBER(G6),G6>0)'], fill=solid(A_BG), font=font(19, True, A_TX),
        stopIfTrue=True))
    ws.conditional_formatting.add("G6", FormulaRule(
        formula=['AND(ISNUMBER(G6),G6<=0)'], fill=solid(G_BG), font=font(19, True, G_TX),
        stopIfTrue=True))

    for r in (8, 12):
        ws.row_dimensions[r].height = 6

    # --------------------------------------------------------- detail pairs --- #
    band(ws, 13, NC, "PROJECT DETAIL (AUTOMATIC)", BLUE)
    pairs = [
        (14, [
            ("SCHEDULE VARIANCE (DAYS)", f'={idx(PRJ_SV)}', "0"),
            ("FORECAST vs APPROVED", f'=IFERROR(INDEX({PRJ_APPROVED},{m})-INDEX({PRJ_FORECAST},{m}),"-")', "#,##0"),
            ("COST UTILISATION", f'=IFERROR(INDEX({PRJ_ACTUAL},{m})/INDEX({PRJ_APPROVED},{m}),"-")', FMT["pct"]),
        ]),
        (15, [
            ("BASELINE END", f'={idx(PRJ_BASE)}', FMT["date"]),
            ("FORECAST END", f'={idx(PRJ_FCEND)}', FMT["date"]),
            ("ACTUAL END", f'={idx(R(PRJ, "Actual End Date"))}', FMT["date"]),
        ]),
        (16, [
            ("OPEN RISKS", f'=COUNTIFS({RAID_PROJ},{SEL},{RAID_TYPE},"Risk",{RAID_STATUS},"<>Closed")', "0"),
            ("OPEN ISSUES", f'=COUNTIFS({RAID_PROJ},{SEL},{RAID_TYPE},"Issue",{RAID_STATUS},"<>Closed")', "0"),
            ("OVERDUE ACTIONS", f'=COUNTIFS({ACT_PROJ},{SEL},{ACT_DOD},">0")', "0"),
        ]),
        (17, [
            ("NEXT MILESTONE", f'=IFERROR(INDEX({MS_NAME},MATCH({SEL}&"|1",{MS_PROJKEY},0)),"-")', None),
            ("MILESTONE DATE", f'=IFERROR(INDEX({MS_FC},MATCH({SEL}&"|1",{MS_PROJKEY},0)),"-")', FMT["date"]),
            ("TASKS COMPLETE",
             f'=IF(COUNTIF({WBS_PROJ},{SEL})=0,"-",COUNTIFS({WBS_PROJ},{SEL},{WBS_STATUS},"Complete")'
             f'&" of "&COUNTIF({WBS_PROJ},{SEL}))', None),
        ]),
    ]
    label_span = [(1, 2, 3, 5), (6, 7, 8, 10), (11, 12, 13, 14)]
    for row, items in pairs:
        for (l1, l2, v1, v2), (label, formula, fmt) in zip(label_span, items):
            merge_style(ws, row, l1, row, l2, label, solid(PALE), font(8.5, True, GRAY_TX),
                        al("right", "center"), border=hair_bottom())
            merge_style(ws, row, v1, row, v2, formula, solid(WHITE), font(10, True, NAVY),
                        al("center", "center"), fmt, border=hair_bottom())
        ws.row_dimensions[row].height = 22

    # --------------------------------------------------------------- panels --- #
    panel_header(ws, 19, 1, 7,
                 f'="TOP OPEN RISKS   ("&COUNTIFS({RAID_PROJ},{SEL},{RAID_TYPE},"Risk",{RAID_STATUS},"<>Closed")&")"')
    panel_header(ws, 19, 8, 14,
                 f'="TOP ISSUES   ("&COUNTIFS({RAID_PROJ},{SEL},{RAID_TYPE},"Issue",{RAID_STATUS},"<>Closed")&")"')
    for i in range(5):
        r = 20 + i
        k = i + 1
        mr = f'MATCH("Risk|"&{SEL}&"|{k}",{RAID_PROJKEY},0)'
        panel_line(ws, r, 1, 7,
                   f'="{TICK}"&IFERROR(INDEX({RAID_DESC},{mr})&"   ("&INDEX({RAID_PRIORITY},{mr})&'
                   f'" - "&INDEX({R(RAID, "Owner")},{mr})&")","")', shade=bool(i % 2))
        mi = f'MATCH("Issue|"&{SEL}&"|{k}",{RAID_PROJKEY},0)'
        panel_line(ws, r, 8, 14,
                   f'="{TICK}"&IFERROR(INDEX({RAID_DESC},{mi})&"   ("&INDEX({RAID_PRIORITY},{mi})&'
                   f'" - "&INDEX({R(RAID, "Owner")},{mi})&")","")', shade=bool(i % 2))

    panel_header(ws, 26, 1, 7,
                 f'="OPEN ACTIONS   ("&COUNTIFS({ACT_PROJ},{SEL},{ACT_STATUS},"<>Complete")&")"')
    panel_header(ws, 26, 8, 14,
                 f'="UPCOMING MILESTONES   ("&COUNTIFS({MS_PROJ},{SEL},{MS_STATUS},"<>Complete")&")"')
    for i in range(5):
        r = 27 + i
        k = i + 1
        ma = f'MATCH({SEL}&"|{k}",{ACT_PROJKEY},0)'
        panel_line(ws, r, 1, 7,
                   f'="{TICK}"&IFERROR(INDEX({ACT_OWNER},{ma})&": "&LEFT(INDEX({ACT_ACTION},{ma}),40)&'
                   f'"   due "&TEXT(INDEX({ACT_DUE},{ma}),"dd-mmm"),"")', shade=bool(i % 2))
        mm = f'MATCH({SEL}&"|{k}",{MS_PROJKEY},0)'
        panel_line(ws, r, 8, 14,
                   f'="{TICK}"&IFERROR(INDEX({MS_NAME},{mm})&"   ("&TEXT(INDEX({MS_FC},{mm}),"dd-mmm-yy")&'
                   f'", "&INDEX({R(MS, "Owner")},{mm})&")","")', shade=bool(i % 2))

    ws.row_dimensions[32].height = 6
    band(ws, 33, NC, "PROJECT CHARTS - RECALCULATE WHEN YOU CHANGE THE PROJECT SELECTED ABOVE", BLUE)

    _chart_data(ws, m)
    _charts(ws)

    ws.freeze_panes = "A4"
    setup_print(ws, NC, 81, title_rows="1:3")
    return ws


def _chart_data(ws, m):
    cd_label(ws, 1, "AUTOMATIC CHART DATA - FORMULAS ONLY, DO NOT EDIT")

    cd_label(ws, 3, "BLOCK 1 - TASK STATUS")
    cd_header(ws, 4, ["Status", "Tasks"])
    for i, s in enumerate(["Not Started", "In Progress", "Complete", "Blocked"]):
        cd_value(ws, 5 + i, 0, s)
        cd_value(ws, 5 + i, 1, f'=COUNTIFS({WBS_PROJ},{SEL},{WBS_STATUS},P{5 + i})', "0")

    cd_label(ws, 10, "BLOCK 2 - TASKS BY OWNER")
    cd_header(ws, 11, ["Owner", "Tasks"])
    for i, s in enumerate(OWNERS):
        cd_value(ws, 12 + i, 0, s)
        cd_value(ws, 12 + i, 1, f'=COUNTIFS({WBS_PROJ},{SEL},{WBS_OWNER},P{12 + i})', "0")

    cd_label(ws, 21, "BLOCK 3 - OPEN RAID BY TYPE")
    cd_header(ws, 22, ["Type", "Open items"])
    for i, s in enumerate(["Risk", "Assumption", "Issue", "Dependency"]):
        cd_value(ws, 23 + i, 0, s)
        cd_value(ws, 23 + i, 1,
                 f'=COUNTIFS({RAID_PROJ},{SEL},{RAID_TYPE},P{23 + i},{RAID_STATUS},"<>Closed")', "0")

    cd_label(ws, 28, "BLOCK 4 - COST BY CATEGORY")
    cd_header(ws, 29, ["Category", "Approved", "Actual", "Forecast"])
    cats = ["Personnel", "Software", "Hardware", "Procurement",
            "Consulting", "Training", "Travel", "Other"]
    for i, s in enumerate(cats):
        r = 30 + i
        cd_value(ws, r, 0, s)
        cd_value(ws, r, 1, f'=SUMIFS({CST_APPROVED},{CST_PROJ},{SEL},{CST_CAT},$P{r})', "#,##0")
        cd_value(ws, r, 2, f'=SUMIFS({CST_ACTUAL},{CST_PROJ},{SEL},{CST_CAT},$P{r})', "#,##0")
        cd_value(ws, r, 3, f'=SUMIFS({CST_FORECAST},{CST_PROJ},{SEL},{CST_CAT},$P{r})', "#,##0")

    cd_label(ws, 39, "BLOCK 5 - MILESTONE STATUS")
    cd_header(ws, 40, ["Status", "Milestones"])
    for i, s in enumerate(["Not Started", "In Progress", "Complete", "Delayed", "Cancelled"]):
        cd_value(ws, 41 + i, 0, s)
        cd_value(ws, 41 + i, 1, f'=COUNTIFS({MS_PROJ},{SEL},{MS_STATUS},P{41 + i})', "0")

    cd_label(ws, 47, "BLOCK 6 - ACTION STATUS")
    cd_header(ws, 48, ["Status", "Actions"])
    for i, s in enumerate(["Open", "In Progress", "Complete", "Cancelled"]):
        cd_value(ws, 49 + i, 0, s)
        cd_value(ws, 49 + i, 1,
                 f'=COUNTIFS({R(ACT, "Project")},{SEL},{ACT_STATUS},P{49 + i})', "0")


def _charts(ws):
    doughnut(ws, "TASK STATUS", href(ws, 16, 5, 16, 8), href(ws, 17, 4, 17, 8),
             ["9AA7B4", "E8A33D", "2E7D32", "C0392B"], "A34")
    bar(ws, "TASKS BY OWNER", href(ws, 16, 12, 16, 19), href(ws, 17, 11, 17, 19),
        "H34", colors="2E75B6", horizontal=True, legend=False)
    bar(ws, "OPEN RAID BY TYPE", href(ws, 16, 23, 16, 26), href(ws, 17, 22, 17, 26),
        "A50", colors="C0392B", legend=False)
    bar(ws, "COST BY CATEGORY", href(ws, 16, 30, 16, 37), href(ws, 17, 29, 19, 37),
        "H50", colors=["2E75B6", "2E7D32", "C0392B"])
    bar(ws, "MILESTONE STATUS", href(ws, 16, 41, 16, 45), href(ws, 17, 40, 17, 45),
        "A66", colors="2E75B6", legend=False,
        point_colors=["9AA7B4", "E8A33D", "2E7D32", "C0392B", "B7C3CF"])
    doughnut(ws, "ACTION STATUS", href(ws, 16, 49, 16, 52), href(ws, 17, 48, 17, 52),
             ["9AA7B4", "E8A33D", "2E7D32", "B7C3CF"], "H66")
