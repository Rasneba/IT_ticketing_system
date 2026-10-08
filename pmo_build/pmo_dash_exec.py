"""01_EXECUTIVE_DASHBOARD - 14 KPI cards, management panels, filter view, 10 charts."""

from openpyxl.styles import Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

from pmo_core import (NAVY, BLUE, BLUE_MID, BLUE_LT, PALE, WHITE, GRAY_TX,
                      G_BG, G_TX, A_BG, A_TX, R_BG, R_TX,
                      font, solid, al, band, sheet_title, add_list_validation,
                      cf_map, cf_databar, setup_print, FMT, STATUS_MAP, RAG_MAP)
from pmo_charts import doughnut, bar, line
from pmo_dash_common import (R, PRJ, MS, RAID, ACT, CHG, DEC, QUAL, CLO,
                             PRJ_ID, PRJ_NAME, PRJ_STATUS, PRJ_RAG, PRJ_MGR, PRJ_DEPT,
                             PRJ_START, PRJ_BASE, PRJ_FCEND, PRJ_PCT, PRJ_APPROVED,
                             PRJ_ACTUAL, PRJ_FORECAST, PRJ_NEXTMS, PRJ_MSDATE,
                             PRJ_TIMELINE, PRJ_BA, PRJ_CHARTIDX, PRJ_REDRANK, PRJ_VIEWRANK,
                             RAID_TYPE, RAID_STATUS, RAID_PRIORITY, RAID_DESC, RAID_PROJ,
                             RAID_RANK, RAID_ESC,
                             ACT_DOD, ACT_RANK, ACT_OWNER, ACT_ACTION,
                             MS_KEY, MS_NAME, MS_PROJ, MS_OWNER, MS_FC, MS_BASE, MS_ACTUAL,
                             MS_STATUS, MS_RAG, MS_DAYS,
                             CHG_STATUS, CHG_RANK, CHG_ID, CHG_DESC, CHG_COST,
                             DEC_STATUS, DEC_DUE,
                             QL_RESULT, CL_ID, CL_STATUS,
                             TICK, merge_style, kpi_card, panel_header, panel_line,
                             cd_label, cd_header, cd_value, href, link_cell, thin_border,
                             hair_bottom)

OWNERS = ["Ahmed Khan", "Lisa Chen", "David Okoro", "Maria Silva",
          "Priya Raman", "Tom Becker", "Grace Alemu", "Sarah Mitchell"]


def build_exec_dashboard(wb):
    ws = wb.create_sheet("01_EXECUTIVE_DASHBOARD")
    ws.sheet_properties.tabColor = NAVY
    ws.sheet_view.showGridLines = False
    NC = 14

    for i in range(1, NC + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    for i in range(16, 21):
        ws.column_dimensions[get_column_letter(i)].width = 17
    ws.column_dimensions["O"].width = 3

    sheet_title(
        ws, "EXECUTIVE DASHBOARD",
        "Live portfolio control panel: 14 KPIs, management attention panels, next milestones, "
        "a filterable portfolio view and 10 charts - every figure is a formula.",
        NC,
        note="All values calculate from the registers (portfolio, RAID, actions, milestones, "
             "changes, decisions). Change a register and this sheet updates automatically.")

    headline = (
        f'="PORTFOLIO: "&A6&" PROJECTS   |   ACTIVE: "&C6&"   |   ON TRACK: "&E6&'
        f'"   |   AT RISK: "&G6&"   |   DELAYED: "&I6&"   |   OPEN RISKS: "&M6&'
        f'"   |   OVERDUE ACTIONS: "&C10&"   |   DECISIONS PENDING: "&M10&'
        f'"   |   AS AT "&TEXT(TODAY(),"DD MMM YYYY")'
    )
    merge_style(ws, 4, 1, 4, NC, headline, solid(NAVY), font(11, True, WHITE),
                al("center", "center"))
    ws.row_dimensions[4].height = 26

    k1 = [
        (1, "TOTAL PROJECTS", f'=COUNTIF({PRJ_NAME},"?*")', "REGISTERED IN PORTFOLIO", "0", None),
        (3, "ACTIVE PROJECTS", f'=COUNTIF({PRJ_STATUS},"Active")',
         '=TEXT(IFERROR(C6/A6,0),"0%")&" OF PORTFOLIO"', "0", None),
        (5, "ON TRACK",
         f'=COUNTIFS({PRJ_RAG},"Green",{PRJ_STATUS},"<>Completed",{PRJ_STATUS},"<>Cancelled")',
         "GREEN RAG, IN FLIGHT", "0", None),
        (7, "AT RISK", f'=COUNTIF({PRJ_RAG},"Yellow")', "YELLOW RAG - WATCH", "0", "amber"),
        (9, "DELAYED", f'=COUNTIF({PRJ_RAG},"Red")', "RED RAG - ACT NOW", "0", "red"),
        (11, "COMPLETED", f'=COUNTIF({PRJ_STATUS},"Completed")', "CLOSED / IN CLOSURE", "0", None),
        (13, "OPEN RISKS", f'=COUNTIFS({RAID_TYPE},"Risk",{RAID_STATUS},"<>Closed")',
         f'="CRITICAL: "&COUNTIFS({RAID_TYPE},"Risk",{RAID_PRIORITY},"Critical",{RAID_STATUS},"<>Closed")',
         "0", "amber"),
    ]
    for col, lab, val, cap, fmt, alert in k1:
        kpi_card(ws, 5, col, lab, val, cap, fmt, alert)

    k2 = [
        (1, "OPEN ISSUES", f'=COUNTIFS({RAID_TYPE},"Issue",{RAID_STATUS},"<>Closed")',
         f'="ESCALATED: "&COUNTIFS({RAID_TYPE},"Issue",{RAID_STATUS},"<>Closed",{RAID_ESC},"Yes")',
         "0", "amber"),
        (3, "OVERDUE ACTIONS", f'=COUNTIF({ACT_DOD},">0")',
         f'="WORST: "&MAX({ACT_DOD})&" DAYS LATE"', "0", "red"),
        (5, "APPROVED AMOUNT", f'=SUM({PRJ_APPROVED})', "TOTAL APPROVED", "#,##0", None),
        (7, "ACTUAL COST", f'=SUM({PRJ_ACTUAL})',
         '=TEXT(IFERROR(G10/E10,0),"0%")&" OF APPROVED"', "#,##0", None),
        (9, "FORECAST COST", f'=SUM({PRJ_FORECAST})',
         '=IF(I10>E10,"OVER APPROVED BY "&TEXT(I10-E10,"#,##0"),"WITHIN APPROVED")', "#,##0", None),
        (13, "PENDING DECISIONS", f'=COUNTIF({DEC_STATUS},"Pending")',
         f'="CHANGES PENDING: "&COUNTIF({CHG_STATUS},"Pending")', "0", "amber"),
    ]
    for col, lab, val, cap, fmt, alert in k2:
        kpi_card(ws, 9, col, lab, val, cap, fmt, alert)

    kpi_card(ws, 9, 11, "MILESTONES DUE 7 DAYS",
             f'=COUNTIFS({MS_FC},">="&TODAY(),{MS_FC},"<="&TODAY()+7,{MS_STATUS},"<>Complete")',
             f'="OVERDUE: "&COUNTIFS({MS_FC},">0",{MS_FC},"<"&TODAY(),{MS_STATUS},"<>Complete")',
             "0", "amber")

    for r in (8, 12):
        ws.row_dimensions[r].height = 6

    # -------------------------------------------------------------- panels --- #
    band(ws, 13, NC, "MANAGEMENT ATTENTION - WHAT NEEDS A DECISION THIS WEEK", BLUE)

    panel_header(ws, 14, 1, 7, f'="RED / DELAYED PROJECTS   ("&COUNTIF({PRJ_RAG},"Red")&")"')
    panel_header(ws, 14, 8, 14,
                 f'="TOP OPEN RISKS BY SCORE   ("&COUNTIFS({RAID_TYPE},"Risk",{RAID_STATUS},"<>Closed")&")"')
    for i in range(5):
        r = 15 + i
        k = i + 1
        m = f'MATCH({k},{PRJ_REDRANK},0)'
        panel_line(ws, r, 1, 7,
                   f'="{TICK}"&IFERROR(INDEX({PRJ_NAME},{m})&"   ["&INDEX({PRJ_STATUS},{m})&"]   "'
                   f'&INDEX({PRJ_MGR},{m})&" - forecast "&TEXT(INDEX({PRJ_FCEND},{m}),"dd-mmm-yy"),"")',
                   shade=bool(i % 2))
        m2 = f'MATCH({k},{RAID_RANK},0)'
        panel_line(ws, r, 8, 14,
                   f'="{TICK}"&IFERROR(INDEX({RAID_DESC},{m2})&"   ("&INDEX({RAID_PRIORITY},{m2})&'
                   f'" - "&INDEX({RAID_PROJ},{m2})&")","")', shade=bool(i % 2))

    panel_header(ws, 21, 1, 7, f'="OVERDUE ACTIONS   ("&COUNTIF({ACT_DOD},">0")&")"')
    panel_header(ws, 21, 8, 14, '="PENDING APPROVALS & DECISIONS"')
    for i in range(5):
        r = 22 + i
        k = i + 1
        ma = f'MATCH({k},{ACT_RANK},0)'
        panel_line(ws, r, 1, 7,
                   f'="{TICK}"&IFERROR(INDEX({ACT_OWNER},{ma})&": "&LEFT(INDEX({ACT_ACTION},{ma}),42)&'
                   f'"  ("&INDEX({ACT_DOD},{ma})&"d late)","")', shade=bool(i % 2))
        if i < 3:
            mc = f'MATCH({k},{CHG_RANK},0)'
            panel_line(ws, r, 8, 14,
                       f'="{TICK}"&IFERROR(INDEX({CHG_ID},{mc})&"   "&LEFT(INDEX({CHG_DESC},{mc}),44)&'
                       f'"   ["&TEXT(INDEX({CHG_COST},{mc}),"#,##0")&"]","")', shade=bool(i % 2))
        elif i == 3:
            panel_line(ws, r, 8, 14,
                       f'="{TICK}"&"Decisions pending: "&COUNTIF({DEC_STATUS},"Pending")&'
                       f'"   |   decisions overdue: "&COUNTIFS({DEC_STATUS},"Pending",{DEC_DUE},"<"&TODAY())&'
                       f'"   |   milestones overdue: "&COUNTIFS({MS_FC},">0",{MS_FC},"<"&TODAY(),{MS_STATUS},"<>Complete")',
                       shade=True)
        else:
            panel_line(ws, r, 8, 14,
                       f'="{TICK}"&"Pending change value: "&TEXT(SUMIFS({CHG_COST},{CHG_STATUS},"Pending"),"#,##0")&'
                       f'"   |   quality failures: "&COUNTIF({QL_RESULT},"Fail")&'
                       f'"   |   closure items open: "&COUNTIFS({CL_STATUS},"<>Complete",{CL_STATUS},"<>N/A",{CL_ID},"?*")',
                       shade=True)

    # ---------------------------------------------------- next 10 milestones - #
    band(ws, 28, NC, "NEXT 10 MILESTONES - AUTOMATIC, CLOSEST FORECAST DATE FIRST", BLUE)
    groups = [(1, 4, "MILESTONE"), (5, 7, "PROJECT"), (8, 9, "OWNER"),
              (10, 11, "FORECAST DATE"), (12, 13, "DAYS"), (14, 14, "RAG")]
    for c1, c2, label in groups:
        merge_style(ws, 29, c1, 29, c2, label, solid(NAVY), font(9, True, WHITE),
                    al("center", "center", wrap=True),
                    border=Border(bottom=Side(style="medium", color=BLUE_MID)))
    ws.row_dimensions[29].height = 22

    for i in range(10):
        r = 30 + i
        k = i + 1
        m = f'MATCH(SMALL({MS_KEY},{k}),{MS_KEY},0)'
        merge_style(ws, r, 1, r, 4, f'=IFERROR(INDEX({MS_NAME},{m}),"")', solid(WHITE),
                    font(9), al("left", "center", indent=1), border=hair_bottom())
        merge_style(ws, r, 5, r, 7, f'=IFERROR(INDEX({MS_PROJ},{m}),"")', solid(WHITE),
                    font(9), al("left", "center", indent=1), border=hair_bottom())
        merge_style(ws, r, 8, r, 9, f'=IFERROR(INDEX({MS_OWNER},{m}),"")', solid(WHITE),
                    font(9), al("left", "center", indent=1), border=hair_bottom())
        merge_style(ws, r, 10, r, 11, f'=IFERROR(INDEX({MS_FC},{m}),"")', solid(WHITE),
                    font(9, True), al("center", "center"), FMT["date"], border=hair_bottom())
        merge_style(ws, r, 12, r, 13, f'=IFERROR(INDEX({MS_DAYS},{m}),"")', solid(WHITE),
                    font(9, True), al("center", "center"), "0", border=hair_bottom())
        merge_style(ws, r, 14, r, 14, f'=IFERROR(INDEX({MS_RAG},{m}),"")', solid(WHITE),
                    font(9, True), al("center", "center"), border=hair_bottom())
        ws.row_dimensions[r].height = 16
    cf_map(ws, "N30:N39", "N", RAG_MAP, first_row=30)
    ws.conditional_formatting.add("L30:M39", FormulaRule(
        formula=['AND(ISNUMBER($L30),$L30<0)'], fill=solid(R_BG), font=font(9, True, R_TX)))
    ws.conditional_formatting.add("L30:M39", FormulaRule(
        formula=['AND(ISNUMBER($L30),$L30>=0,$L30<=7)'], fill=solid(A_BG), font=font(9, True, A_TX)))

    # ------------------------------------------------------ filter view ------- #
    band(ws, 41, NC,
         "PORTFOLIO FILTER VIEW -  PROJECT  |  STATUS  |  RAG  |  MANAGER / DEPARTMENT  |  "
         "TIMELINE  |  PROGRESS  |  BUDGET / ACTUAL  |  NEXT MILESTONE", BLUE)

    bd = thin_border()
    merge_style(ws, 42, 1, 42, 2, "FILTER BY:", solid(NAVY), font(9, True, WHITE),
                al("right", "center"), border=bd)
    for col, dv_name, prompt in ((3, "L_FilterStatus", "Filter by project status"),
                                 (6, "L_FilterRAG", "Filter by RAG"),
                                 (9, "L_FilterMgr", "Filter by project manager")):
        c = ws.cell(row=42, column=col, value="All")
        c.font, c.fill, c.alignment, c.border = font(10, True, NAVY), solid(WHITE), al("center", "center"), bd
        add_list_validation(ws, get_column_letter(col), dv_name, 42, 42, prompt=prompt)
    merge_style(ws, 42, 4, 42, 5, "RAG:", solid(NAVY), font(9, True, WHITE),
                al("right", "center"), border=bd)
    merge_style(ws, 42, 7, 42, 8, "MANAGER:", solid(NAVY), font(9, True, WHITE),
                al("right", "center"), border=bd)
    merge_style(ws, 42, 10, 42, 12,
                f'=COUNTIF({PRJ_VIEWRANK},">0")&" OF "&COUNTIF({PRJ_NAME},"?*")&" PROJECTS SHOWN"',
                solid(PALE), font(9, True, NAVY), al("center", "center"), border=bd)
    link_cell(ws, 42, 13, 14, '=HYPERLINK("#\'02_PROJECT_PORTFOLIO\'!A1","OPEN FULL PORTFOLIO")',
              solid(BLUE_LT))
    ws.row_dimensions[42].height = 24

    view_groups = [(1, 2, "PROJECT"), (3, 3, "STATUS"), (4, 4, "RAG"),
                   (5, 6, "MANAGER / DEPARTMENT"), (7, 8, "TIMELINE"), (9, 9, "PROGRESS"),
                   (10, 11, "BUDGET / ACTUAL"), (12, 14, "NEXT MILESTONE")]
    for c1, c2, label in view_groups:
        merge_style(ws, 43, c1, 43, c2, label, solid(NAVY), font(9, True, WHITE),
                    al("center", "center", wrap=True),
                    border=Border(bottom=Side(style="medium", color=BLUE_MID)))
    ws.row_dimensions[43].height = 26

    for i in range(15):
        r = 44 + i
        k = i + 1
        m = f'MATCH({k},{PRJ_VIEWRANK},0)'
        merge_style(ws, r, 1, r, 2, f'=IFERROR(INDEX({PRJ_ID},{m})&"   "&INDEX({PRJ_NAME},{m}),"")',
                    solid(WHITE), font(9), al("left", "center", indent=1), border=hair_bottom())
        merge_style(ws, r, 3, r, 3, f'=IFERROR(INDEX({PRJ_STATUS},{m}),"")',
                    solid(WHITE), font(9), al("center", "center"), border=hair_bottom())
        merge_style(ws, r, 4, r, 4, f'=IFERROR(INDEX({PRJ_RAG},{m}),"")',
                    solid(WHITE), font(9, True), al("center", "center"), border=hair_bottom())
        merge_style(ws, r, 5, r, 6,
                    f'=IFERROR(INDEX({PRJ_MGR},{m})&" / "&INDEX({PRJ_DEPT},{m}),"")',
                    solid(WHITE), font(9), al("left", "center", indent=1), border=hair_bottom())
        merge_style(ws, r, 7, r, 8, f'=IFERROR(INDEX({PRJ_TIMELINE},{m}),"")',
                    solid(WHITE), font(9), al("left", "center", indent=1), border=hair_bottom())
        merge_style(ws, r, 9, r, 9, f'=IFERROR(INDEX({PRJ_PCT},{m}),"")',
                    solid(WHITE), font(9), al("center", "center"), FMT["pct"], border=hair_bottom())
        merge_style(ws, r, 10, r, 11, f'=IFERROR(INDEX({PRJ_BA},{m}),"")',
                    solid(WHITE), font(9), al("center", "center"), border=hair_bottom())
        merge_style(ws, r, 12, r, 14,
                    f'=IFERROR(INDEX({PRJ_NEXTMS},{m})&IF(INDEX({PRJ_MSDATE},{m})="","",'
                    f'"   ("&TEXT(INDEX({PRJ_MSDATE},{m}),"dd-mmm-yy")&")"),"")',
                    solid(WHITE), font(9), al("left", "center", indent=1), border=hair_bottom())
        ws.row_dimensions[r].height = 16
    cf_map(ws, "C44:C58", "C", STATUS_MAP, first_row=44)
    cf_map(ws, "D44:D58", "D", RAG_MAP, first_row=44)
    cf_databar(ws, "I44:I58")

    ws.row_dimensions[59].height = 6
    band(ws, 60, NC, "PORTFOLIO CHARTS - ALL FORMULA DRIVEN, NO MANUAL UPDATING", BLUE)

    _chart_data(ws)
    _charts(ws)

    ws.freeze_panes = "A4"
    setup_print(ws, NC, 141, title_rows="1:3")
    return ws


def _chart_data(ws):
    cd_label(ws, 1, "AUTOMATIC CHART DATA - FORMULAS ONLY, DO NOT EDIT")

    cd_label(ws, 3, "BLOCK 1 - PROJECT STATUS")
    cd_header(ws, 4, ["Status", "Projects"])
    for i, s in enumerate(["Not Started", "Active", "On Hold", "Completed", "Cancelled"]):
        cd_value(ws, 5 + i, 0, s)
        cd_value(ws, 5 + i, 1, f'=COUNTIF({PRJ_STATUS},P{5 + i})', "0")

    cd_label(ws, 11, "BLOCK 2 - PROJECT RAG")
    cd_header(ws, 12, ["RAG", "Projects"])
    for i, s in enumerate(["Green", "Yellow", "Red"]):
        cd_value(ws, 13 + i, 0, s)
        cd_value(ws, 13 + i, 1, f'=COUNTIF({PRJ_RAG},P{13 + i})', "0")

    cd_label(ws, 17, "BLOCK 3 - PROJECT COMPLETION")
    cd_header(ws, 18, ["Project", "% Complete"])
    for i in range(10):
        r = 19 + i
        cd_value(ws, r, 0, f'=IFERROR(INDEX({PRJ_NAME},MATCH({i + 1},{PRJ_CHARTIDX},0)),"")')
        cd_value(ws, r, 1, f'=IF($P{r}="","",INDEX({PRJ_PCT},MATCH($P{r},{PRJ_NAME},0)))', "0%")

    cd_label(ws, 30, "BLOCK 4 - PLANNED vs ACTUAL PROGRESS")
    cd_header(ws, 31, ["Project", "Planned to date", "Actual"])
    for i in range(10):
        r = 32 + i
        cd_value(ws, r, 0, f'=IFERROR(INDEX({PRJ_NAME},MATCH({i + 1},{PRJ_CHARTIDX},0)),"")')
        m = f'MATCH($P{r},{PRJ_NAME},0)'
        cd_value(ws, r, 1,
                 f'=IF($P{r}="","",IFERROR(MAX(0,MIN(1,(TODAY()-INDEX({PRJ_START},{m}))/'
                 f'(INDEX({PRJ_BASE},{m})-INDEX({PRJ_START},{m})))),""))', "0%")
        cd_value(ws, r, 2, f'=IF($P{r}="","",INDEX({PRJ_PCT},{m}))', "0%")

    cd_label(ws, 43, "BLOCK 5 - APPROVED / ACTUAL / FORECAST")
    cd_header(ws, 44, ["Project", "Approved", "Actual", "Forecast"])
    for i in range(10):
        r = 45 + i
        cd_value(ws, r, 0, f'=IFERROR(INDEX({PRJ_NAME},MATCH({i + 1},{PRJ_CHARTIDX},0)),"")')
        m = f'MATCH($P{r},{PRJ_NAME},0)'
        cd_value(ws, r, 1, f'=IF($P{r}="","",INDEX({PRJ_APPROVED},{m}))', "#,##0")
        cd_value(ws, r, 2, f'=IF($P{r}="","",INDEX({PRJ_ACTUAL},{m}))', "#,##0")
        cd_value(ws, r, 3, f'=IF($P{r}="","",INDEX({PRJ_FORECAST},{m}))', "#,##0")

    cd_label(ws, 56, "BLOCK 6 - MONTHLY MILESTONE DELIVERY (BASELINE vs ACTUAL)")
    cd_header(ws, 57, ["Month", "Month start", "Baseline", "Actual"])
    for i in range(6):
        r = 58 + i
        if i == 0:
            cd_value(ws, r, 1, "=EOMONTH(TODAY(),-1)+1", "mmm-yy")
        else:
            cd_value(ws, r, 1, f"=EDATE(Q{r - 1},1)", "mmm-yy")
        cd_value(ws, r, 0, f'=TEXT($Q{r},"mmm-yy")')
        cd_value(ws, r, 2, f'=COUNTIFS({MS_BASE},">="&$Q{r},{MS_BASE},"<"&EDATE($Q{r},1))', "0")
        cd_value(ws, r, 3, f'=COUNTIFS({MS_ACTUAL},">="&$Q{r},{MS_ACTUAL},"<"&EDATE($Q{r},1))', "0")

    cd_label(ws, 65, "BLOCK 7 - OPEN RISKS BY SEVERITY")
    cd_header(ws, 66, ["Priority", "Open risks"])
    for i, s in enumerate(["Low", "Medium", "High", "Critical"]):
        cd_value(ws, 67 + i, 0, s)
        cd_value(ws, 67 + i, 1,
                 f'=COUNTIFS({RAID_TYPE},"Risk",{RAID_PRIORITY},P{67 + i},{RAID_STATUS},"<>Closed")', "0")

    cd_label(ws, 72, "BLOCK 8 - ISSUES BY STATUS")
    cd_header(ws, 73, ["Status", "Issues"])
    for i, s in enumerate(["Open", "In Progress", "Closed"]):
        cd_value(ws, 74 + i, 0, s)
        cd_value(ws, 74 + i, 1, f'=COUNTIFS({RAID_TYPE},"Issue",{RAID_STATUS},P{74 + i})', "0")

    cd_label(ws, 78, "BLOCK 9 - OVERDUE ACTIONS BY OWNER")
    cd_header(ws, 79, ["Owner", "Overdue actions"])
    for i, s in enumerate(OWNERS):
        cd_value(ws, 80 + i, 0, s)
        cd_value(ws, 80 + i, 1, f'=COUNTIFS({ACT_OWNER},P{80 + i},{ACT_DOD},">0")', "0")

    cd_label(ws, 89, "BLOCK 10 - UPCOMING MILESTONES (NEXT 30 DAYS)")
    cd_header(ws, 90, ["Window", "Milestones"])
    cd_value(ws, 91, 0, "Next 7 days")
    cd_value(ws, 91, 1,
             f'=COUNTIFS({MS_FC},">="&TODAY(),{MS_FC},"<="&TODAY()+7,{MS_STATUS},"<>Complete")', "0")
    cd_value(ws, 92, 0, "8 - 14 days")
    cd_value(ws, 92, 1,
             f'=COUNTIFS({MS_FC},">"&TODAY()+7,{MS_FC},"<="&TODAY()+14,{MS_STATUS},"<>Complete")', "0")
    cd_value(ws, 93, 0, "15 - 30 days")
    cd_value(ws, 93, 1,
             f'=COUNTIFS({MS_FC},">"&TODAY()+14,{MS_FC},"<="&TODAY()+30,{MS_STATUS},"<>Complete")', "0")


def _charts(ws):
    doughnut(ws, "PROJECT STATUS", href(ws, 16, 5, 16, 9), href(ws, 17, 4, 17, 9),
             ["9AA7B4", "2E75B6", "E8A33D", "2E7D32", "C0392B"], "A61")
    doughnut(ws, "PROJECT RAG", href(ws, 16, 13, 16, 15), href(ws, 17, 12, 17, 15),
             ["2E7D32", "E8A33D", "C0392B"], "H61")
    bar(ws, "PROJECT COMPLETION", href(ws, 16, 19, 16, 28), href(ws, 17, 18, 17, 28),
        "A77", colors="2E75B6", horizontal=True, legend=False, num_fmt="0%")
    bar(ws, "PLANNED vs ACTUAL PROGRESS", href(ws, 16, 32, 16, 41), href(ws, 17, 31, 18, 41),
        "H77", colors=["B7C3CF", "2E75B6"])
    bar(ws, "APPROVED vs ACTUAL vs FORECAST", href(ws, 16, 45, 16, 54), href(ws, 17, 44, 19, 54),
        "A93", colors=["2E75B6", "2E7D32", "C0392B"])
    line(ws, "MILESTONES DELIVERED PER MONTH", href(ws, 16, 58, 16, 63), href(ws, 18, 57, 19, 63),
         "H93", colors=["2E75B6", "2E7D32"])
    bar(ws, "OPEN RISKS BY SEVERITY", href(ws, 16, 67, 16, 70), href(ws, 17, 66, 17, 70),
        "A109", legend=False, point_colors=["2E7D32", "E8A33D", "C0392B", "8E2C22"])
    bar(ws, "ISSUES BY STATUS", href(ws, 16, 74, 16, 76), href(ws, 17, 73, 17, 76),
        "H109", legend=False, point_colors=["E8A33D", "2E75B6", "2E7D32"])
    bar(ws, "OVERDUE ACTIONS BY OWNER", href(ws, 16, 80, 16, 87), href(ws, 17, 79, 17, 87),
        "A125", colors="C0392B", horizontal=True, legend=False)
    bar(ws, "MILESTONES DUE - NEXT 30 DAYS", href(ws, 16, 91, 16, 93), href(ws, 17, 90, 17, 93),
        "H125", colors="2E75B6", legend=False)
